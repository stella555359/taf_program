#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Fix kpi_custom_rules.xlsx - Restore data validation dropdowns and sheet formatting
Updated: ivy0203 - Use new core strategies
"""

from pathlib import Path
from datetime import datetime
import json
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Border, Side

# Core strategies (ivy0203 updated)
STRATEGIES = ['STANDARD_SPC', 'ABS_THRESHOLD', 'IQR', 'POISSON', 'RELATIVE', 
              'Z_SCORE', 'KS_TEST', 'SKIP']

# Legacy strategy mapping (for updating Default_Strategy column)
LEGACY_STRATEGY_MAP = {
    'RANGE_CONTEXT': 'IQR',
    'RELATIVE_CHANGE': 'RELATIVE',
    'ABS_DELTA_TREND': 'Z_SCORE',
    'QUANTILES': 'IQR',
    'TAIL_AWARE': 'IQR',
    'WINDOW_COUNT': 'POISSON',
    'SEMANTIC_RULE': 'SKIP',
}

# Styles
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
edit_fill = PatternFill(start_color="FFFFD7", end_color="FFFFD7", fill_type="solid")  # Yellow
degraded_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Red
improved_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
attention_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Orange
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


def fix_excel_validation():
    base_dir = Path(__file__).parent.parent
    excel_path = base_dir / 'lib' / 'kpi_xml' / 'kpi_custom_rules.xlsx'
    history_file = base_dir / 'lib' / 'kpi_xml' / 'kpi_rules_history.json'
    
    print(f"Fixing: {excel_path}")
    
    # Load existing file
    wb = load_workbook(excel_path)
    
    # ========== Fix Sheet 1: Custom Rules ==========
    ws = wb['1_Custom_Rules']
    
    # Remove existing validations if any
    ws.data_validations.dataValidation = []
    
    # Create data validations
    strategies_str = ','.join(STRATEGIES)
    dv_strategy = DataValidation(type='list', formula1=f'"{strategies_str}"', allow_blank=True)
    dv_strategy.error = '请从下拉列表中选择策略'
    dv_strategy.errorTitle = '无效的策略'
    
    dv_enabled = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
    
    # Add to sheet
    ws.add_data_validation(dv_strategy)
    ws.add_data_validation(dv_enabled)
    
    # Get max row
    max_row = ws.max_row
    print(f'Total rows in 1_Custom_Rules: {max_row}')
    
    # Apply validations and styles to columns E (Override_Strategy) and F (Enabled)
    # Also update Default_Strategy (column D) to use new strategy names
    degraded_count = 0
    improved_count = 0
    legacy_updated_count = 0
    
    for row in range(2, max_row + 1):
        # Default_Strategy - column D - convert legacy strategies to core
        cell_d = ws.cell(row=row, column=4)
        old_strategy = str(cell_d.value or '').strip()
        if old_strategy in LEGACY_STRATEGY_MAP:
            cell_d.value = LEGACY_STRATEGY_MAP[old_strategy]
            legacy_updated_count += 1
        
        # Override_Strategy - column E - also convert legacy strategies
        cell_e = ws.cell(row=row, column=5)
        override_strategy = str(cell_e.value or '').strip()
        if override_strategy in LEGACY_STRATEGY_MAP:
            cell_e.value = LEGACY_STRATEGY_MAP[override_strategy]
        cell_e.fill = edit_fill
        dv_strategy.add(cell_e)
        
        # Enabled - column F, also convert boolean to string
        cell_f = ws.cell(row=row, column=6)
        if cell_f.value is True:
            cell_f.value = 'TRUE'
        elif cell_f.value is False:
            cell_f.value = 'FALSE'
        cell_f.fill = edit_fill
        dv_enabled.add(cell_f)
        
        # Change_Type - column G - apply colors
        cell_g = ws.cell(row=row, column=7)
        change_type = str(cell_g.value or '')
        if 'Other→Regression' in change_type or '→Regression' in change_type:
            cell_g.fill = degraded_fill
            degraded_count += 1
        elif 'Regression→Other' in change_type or '→Other' in change_type:
            cell_g.fill = improved_fill
            improved_count += 1
        
        # Priority - column H - apply attention color for HIGH
        cell_h = ws.cell(row=row, column=8)
        if cell_h.value and 'HIGH' in str(cell_h.value):
            cell_h.fill = attention_fill
    
    print(f'  Applied strategy dropdown to E2:E{max_row}')
    print(f'  Applied enabled dropdown to F2:F{max_row}')
    print(f'  Updated {legacy_updated_count} legacy strategies to new names')
    print(f'  Degraded (red): {degraded_count}, Improved (green): {improved_count}')
    
    # ========== Fix Sheet 3: Quick Summary ==========
    if '3_Quick_Summary' in wb.sheetnames:
        ws_summary = wb['3_Quick_Summary']
        
        # Count rules
        user_rules_count = sum(1 for row in range(2, max_row + 1) 
                              if ws.cell(row=row, column=6).value == 'TRUE')
        
        # Update summary data
        summary_data = [
            "",
            f"📊 数据统计 (更新时间: {datetime.now().strftime('%Y-%m-%d %H:%M')})",
            f"  总KPI数量: {max_row - 1}",
            f"  Other→Regression (需关注): {degraded_count} 个",
            f"  Regression→Other (已改善): {improved_count} 个",
            f"  用户已配置规则: {user_rules_count} 个",
            "",
            "🔧 推荐操作步骤:",
            "  1. 点击 Sheet '1_Custom_Rules'",
            "  2. 点击 'Change_Type' 列的筛选按钮",
            "  3. 选择 'Other→Regression ⚠️' 进行筛选",
            f"  4. 查看这 {degraded_count} 个KPI，决定是否需要修改策略",
            "  5. 如需修改：在 Override_Strategy 选择新策略，Enabled 设为 TRUE",
            "  6. 保存文件",
            "  7. 运行: python scripts/manage_custom_rules.py --apply",
            "",
            "💡 提示:",
            "  - 红色背景 = Other→Regression (可能需要改回 STANDARD_SPC)",
            "  - 绿色背景 = Regression→Other (新策略生效，通常不需要修改)",
            "  - 黄色背景 = 可编辑列",
            "  - 大部分情况下，只需关注红色行",
        ]
        
        # Clear and rewrite
        for row in ws_summary.iter_rows():
            for cell in row:
                cell.value = None
        
        ws_summary['A1'] = "快速操作指南"
        ws_summary['A1'].font = Font(bold=True, size=14)
        
        for r_idx, text in enumerate(summary_data, 3):
            cell = ws_summary.cell(row=r_idx, column=1, value=text)
            if text.startswith("📊") or text.startswith("🔧") or text.startswith("💡"):
                cell.font = Font(bold=True)
        
        ws_summary.column_dimensions['A'].width = 80
        print('  Updated Sheet 3_Quick_Summary')
    
    # ========== Fix Sheet 4: Modification History ==========
    if '4_Modification_History' in wb.sheetnames and history_file.exists():
        ws_history = wb['4_Modification_History']
        
        # Clear existing data (keep headers)
        for row in range(2, ws_history.max_row + 1):
            for col in range(1, 7):
                ws_history.cell(row=row, column=col).value = None
        
        # Reload from history file
        try:
            with open(history_file, 'r', encoding='utf-8') as f:
                history = json.load(f)
            
            # Write last 100 records
            for h_idx, record in enumerate(history[-100:], 2):
                ws_history.cell(row=h_idx, column=1, value=record.get('timestamp', '')).border = thin_border
                ws_history.cell(row=h_idx, column=2, value=record.get('code', '')).border = thin_border
                ws_history.cell(row=h_idx, column=3, value=record.get('kpi_name', '')).border = thin_border
                ws_history.cell(row=h_idx, column=4, value=record.get('original_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=5, value=record.get('new_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=6, value=record.get('modified_by', 'user')).border = thin_border
            
            print(f'  Updated Sheet 4_Modification_History with {min(len(history), 100)} records')
        except Exception as e:
            print(f'  Warning: Could not load history: {e}')
    
    # Save
    wb.save(excel_path)
    print('\n✅ Excel file fixed successfully!')
    print('Now you can use dropdown lists in Excel to modify rules.')


if __name__ == "__main__":
    fix_excel_validation()
