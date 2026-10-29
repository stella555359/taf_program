#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Compare OLD and NEW HTML reports for KPI detection changes"""

import re
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def extract_kpi_data(html_content):
    """Extract kpiData JSON from HTML"""
    match = re.search(r'const kpiData = ({.*?});', html_content, re.DOTALL)
    if match:
        return json.loads(match.group(1))
    return {}

def main():
    base_dir = Path(__file__).parent.parent
    
    old_path = base_dir / 'reports/0126_bka/0123.2_26R2.0123.7_T076.SCF.26R1psot5_32UEs_httpDL.bigSmall2_1csv9run_trend_viewer.html'
    new_path = base_dir / 'reports/0123.2_26R2.0123.7_T076.SCF.26R1psot5_32UEs_httpDL.bigSmall2_1csv9run_trend_viewer.html'
    
    with open(old_path, 'r', encoding='utf-8') as f:
        old_html = f.read()
    with open(new_path, 'r', encoding='utf-8') as f:
        new_html = f.read()
    
    old_data = extract_kpi_data(old_html)
    new_data = extract_kpi_data(new_html)
    
    # Get regression sets
    old_regression = set(k for k, v in old_data.items() if v.get('severity') == 'Regression')
    new_regression = set(k for k, v in new_data.items() if v.get('severity') == 'Regression')
    
    to_other = old_regression - new_regression  # Regression -> Other
    to_regression = new_regression - old_regression  # Other -> Regression
    
    # Build DataFrames for Excel
    rows_improved = []
    for code in sorted(to_other):
        o = old_data.get(code, {})
        n = new_data.get(code, {})
        rows_improved.append({
            'Code': code,
            'KPI_Name': n.get('kpi_name', ''),
            'OLD_Severity': o.get('severity', ''),
            'NEW_Severity': n.get('severity', ''),
            'KPI_Type': n.get('kpi_type', ''),
            'Detection_Strategy': n.get('detection_strategy', ''),
            'OLD_Sigma': o.get('sigma_level'),
            'NEW_Sigma': n.get('sigma_level'),
            'Anomaly_Count': n.get('anomaly_count', 0),
            'Total_Points': n.get('total_points', 0),
            'Units': n.get('units', ''),
            'Related_Counters': ', '.join(n.get('related_counters', [])[:5]) if n.get('related_counters') else '',
        })
    
    rows_degraded = []
    for code in sorted(to_regression):
        o = old_data.get(code, {})
        n = new_data.get(code, {})
        rows_degraded.append({
            'Code': code,
            'KPI_Name': n.get('kpi_name', ''),
            'OLD_Severity': o.get('severity', ''),
            'NEW_Severity': n.get('severity', ''),
            'KPI_Type': n.get('kpi_type', ''),
            'Detection_Strategy': n.get('detection_strategy', ''),
            'OLD_Sigma': o.get('sigma_level'),
            'NEW_Sigma': n.get('sigma_level'),
            'Anomaly_Count': n.get('anomaly_count', 0),
            'Total_Points': n.get('total_points', 0),
            'Units': n.get('units', ''),
            'Related_Counters': ', '.join(n.get('related_counters', [])[:5]) if n.get('related_counters') else '',
        })
    
    df_improved = pd.DataFrame(rows_improved)
    df_degraded = pd.DataFrame(rows_degraded)
    
    # Create Excel workbook
    output_path = base_dir / 'reports' / 'kpi_change_comparison.xlsx'
    
    wb = Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    improved_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    degraded_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "1_Summary"
    
    def count_severity(data):
        counts = {}
        for k, v in data.items():
            sev = v.get('severity', 'Unknown')
            counts[sev] = counts.get(sev, 0) + 1
        return counts
    
    old_counts = count_severity(old_data)
    new_counts = count_severity(new_data)
    
    ws_summary['A1'] = "KPI Severity Comparison Summary"
    ws_summary['A1'].font = Font(bold=True, size=14)
    ws_summary.merge_cells('A1:E1')
    
    headers = ['Severity', 'OLD Count', 'NEW Count', 'Change', 'Description']
    for col, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=3, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    for row, sev in enumerate(['Regression', 'Suspicious', 'Normal', 'NoHistory'], 4):
        old_c = old_counts.get(sev, 0)
        new_c = new_counts.get(sev, 0)
        diff = new_c - old_c
        desc = ''
        if sev == 'Regression':
            desc = f'减少{-diff}个' if diff < 0 else f'增加{diff}个'
        
        ws_summary.cell(row=row, column=1, value=sev).border = thin_border
        ws_summary.cell(row=row, column=2, value=old_c).border = thin_border
        ws_summary.cell(row=row, column=3, value=new_c).border = thin_border
        change_cell = ws_summary.cell(row=row, column=4, value=diff)
        change_cell.border = thin_border
        if diff < 0:
            change_cell.fill = improved_fill
        elif diff > 0:
            change_cell.fill = degraded_fill
        ws_summary.cell(row=row, column=5, value=desc).border = thin_border
    
    # Change summary
    ws_summary['A9'] = "Change Summary"
    ws_summary['A9'].font = Font(bold=True, size=12)
    ws_summary['A10'] = f"Regression → Other (改善): {len(to_other)}个"
    ws_summary['A11'] = f"Other → Regression (恶化): {len(to_regression)}个"
    ws_summary['A12'] = f"净变化: {len(to_regression) - len(to_other)}个"
    
    # Type breakdown for improved
    ws_summary['A14'] = "Regression → Other 按KPI类型统计"
    ws_summary['A14'].font = Font(bold=True)
    type_stats = {}
    for code in to_other:
        kpi_type = new_data[code].get('kpi_type', 'unknown')
        new_sev = new_data[code].get('severity')
        key = f"{kpi_type} → {new_sev}"
        type_stats[key] = type_stats.get(key, 0) + 1
    
    row = 15
    for k, v in sorted(type_stats.items(), key=lambda x: -x[1]):
        ws_summary.cell(row=row, column=1, value=k)
        ws_summary.cell(row=row, column=2, value=v)
        row += 1
    
    # Type breakdown for degraded
    ws_summary.cell(row=row+1, column=1, value="Other → Regression 按KPI类型统计").font = Font(bold=True)
    type_stats2 = {}
    for code in to_regression:
        kpi_type = new_data[code].get('kpi_type', 'unknown')
        old_sev = old_data[code].get('severity')
        key = f"{old_sev} → Regression ({kpi_type})"
        type_stats2[key] = type_stats2.get(key, 0) + 1
    
    row += 2
    for k, v in sorted(type_stats2.items(), key=lambda x: -x[1]):
        ws_summary.cell(row=row, column=1, value=k)
        ws_summary.cell(row=row, column=2, value=v)
        row += 1
    
    ws_summary.column_dimensions['A'].width = 40
    ws_summary.column_dimensions['B'].width = 12
    ws_summary.column_dimensions['C'].width = 12
    ws_summary.column_dimensions['D'].width = 10
    ws_summary.column_dimensions['E'].width = 20
    
    # Sheet 2: Regression -> Other (Improved)
    ws_improved = wb.create_sheet(title="2_Regression_to_Other")
    
    if not df_improved.empty:
        for col, header in enumerate(df_improved.columns, 1):
            cell = ws_improved.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        
        for row_idx, row_data in enumerate(df_improved.itertuples(index=False), 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws_improved.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx == 4:  # NEW_Severity
                    if value == 'Normal':
                        cell.fill = improved_fill
    
    # Set column widths
    col_widths = [18, 50, 12, 12, 18, 20, 10, 10, 12, 12, 10, 40]
    for i, w in enumerate(col_widths, 1):
        ws_improved.column_dimensions[chr(64+i)].width = w
    
    ws_improved.freeze_panes = 'A2'
    
    # Sheet 3: Other -> Regression (Degraded)
    ws_degraded = wb.create_sheet(title="3_Other_to_Regression")
    
    if not df_degraded.empty:
        for col, header in enumerate(df_degraded.columns, 1):
            cell = ws_degraded.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        
        for row_idx, row_data in enumerate(df_degraded.itertuples(index=False), 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws_degraded.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx == 4:  # NEW_Severity
                    cell.fill = degraded_fill
    
    for i, w in enumerate(col_widths, 1):
        ws_degraded.column_dimensions[chr(64+i)].width = w
    
    ws_degraded.freeze_panes = 'A2'
    
    # Sheet 4: 4 Specific KPIs
    ws_specific = wb.create_sheet(title="4_Specified_KPIs")
    
    specific_codes = ['NR_5351a', 'NR_6090a', 'NR_6094a', 'NR_6023a']
    headers = ['Code', 'KPI_Name', 'OLD_Severity', 'NEW_Severity', 'KPI_Type', 'Strategy', 
               'OLD_Sigma', 'NEW_Sigma', 'Anomaly', 'Units', 'Related_Counters', 'Analysis']
    
    for col, h in enumerate(headers, 1):
        cell = ws_specific.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    analysis_map = {
        'NR_5351a': '下限饱和型KPI(floor_saturated)，使用ABS_THRESHOLD判定，不再误报为Regression',
        'NR_6090a': '比例型KPI(ratio_type)，使用RELATIVE_CHANGE判定，σ=2.26不到3σ阈值',
        'NR_6094a': '延迟型KPI(latency_type)，使用TAIL_AWARE分位数判定',
        'NR_6023a': '配置型KPI(config_type)，SEMANTIC_RULE检测到配置变化，升级为Regression'
    }
    
    for row_idx, code in enumerate(specific_codes, 2):
        o = old_data.get(code, {})
        n = new_data.get(code, {})
        
        values = [
            code,
            n.get('kpi_name', ''),
            o.get('severity', ''),
            n.get('severity', ''),
            n.get('kpi_type', ''),
            n.get('detection_strategy', ''),
            o.get('sigma_level'),
            n.get('sigma_level'),
            f"{n.get('anomaly_count', 0)}/{n.get('total_points', 0)}",
            n.get('units', ''),
            ', '.join(n.get('related_counters', [])[:3]) if n.get('related_counters') else '',
            analysis_map.get(code, '')
        ]
        
        for col_idx, value in enumerate(values, 1):
            cell = ws_specific.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx == 4:
                if value == 'Normal':
                    cell.fill = improved_fill
                elif value == 'Regression':
                    cell.fill = degraded_fill
    
    col_widths_spec = [12, 50, 12, 12, 18, 18, 10, 10, 10, 8, 30, 60]
    for i, w in enumerate(col_widths_spec, 1):
        ws_specific.column_dimensions[chr(64+i)].width = w
    
    # Save
    wb.save(output_path)
    print(f"Excel report saved to: {output_path}")
    print(f"\nSheets:")
    print(f"  1_Summary: 总览统计")
    print(f"  2_Regression_to_Other: {len(to_other)}个改善的KPI")
    print(f"  3_Other_to_Regression: {len(to_regression)}个恶化的KPI")
    print(f"  4_Specified_KPIs: 4个指定KPI详情")
    
    # Generate custom rules Excel file
    generate_custom_rules_excel(base_dir, to_regression, new_data)


def generate_custom_rules_excel(base_dir: Path, to_regression: set, new_data: dict):
    """Generate an Excel file for custom rule editing"""
    from openpyxl.worksheet.datavalidation import DataValidation
    
    STRATEGIES = [
        "STANDARD_SPC",
        "ABS_THRESHOLD",
        "RANGE_CONTEXT",
        "RELATIVE_CHANGE",
        "ABS_DELTA_TREND",
        "QUANTILES",
        "TAIL_AWARE",
        "WINDOW_COUNT",
        "SEMANTIC_RULE",
        "(使用默认)"
    ]
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    edit_fill = PatternFill(start_color="FFFFD7", end_color="FFFFD7", fill_type="solid")  # Light yellow
    improved_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    degraded_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    excel_path = base_dir / 'lib' / 'kpi_xml' / 'kpi_custom_rules.xlsx'
    history_json = base_dir / 'lib' / 'kpi_xml' / 'kpi_rules_history.json'
    
    # Load existing Excel if exists (to preserve user edits)
    existing_overrides = {}
    if excel_path.exists():
        try:
            df_existing = pd.read_excel(excel_path, sheet_name='1_Custom_Rules')
            for _, row in df_existing.iterrows():
                code = str(row.get('Code', '')).strip()
                if code:
                    existing_overrides[code] = {
                        'Override_Strategy': str(row.get('Override_Strategy', '')).strip(),
                        'Enabled': str(row.get('Enabled', 'FALSE')).strip(),
                        'Comment': str(row.get('Comment', '')).strip()
                    }
        except Exception as e:
            print(f"Warning: Could not read existing Excel: {e}")
    
    wb = Workbook()
    
    # Sheet 1: Custom Rules (editable)
    ws_rules = wb.active
    ws_rules.title = "1_Custom_Rules"
    
    headers = ['Code', 'KPI_Name', 'KPI_Type', 'Default_Strategy', 'Override_Strategy', 'Enabled', 'Change_Type', 'Comment']
    for col, h in enumerate(headers, 1):
        cell = ws_rules.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    # Data validation for Override_Strategy
    dv_strategy = DataValidation(
        type="list",
        formula1=f'"{",".join(STRATEGIES)}"',
        allow_blank=True
    )
    dv_strategy.error = "请从下拉列表中选择策略"
    ws_rules.add_data_validation(dv_strategy)
    
    # Data validation for Enabled
    dv_enabled = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws_rules.add_data_validation(dv_enabled)
    
    # Add all codes from new_data
    row_idx = 2
    all_codes = set(new_data.keys())
    
    for code in sorted(all_codes):
        info = new_data.get(code, {})
        
        # Determine change type
        if code in to_regression:
            change_type = "Other→Regression"
        else:
            change_type = ""
        
        # Check for existing overrides
        existing = existing_overrides.get(code, {})
        override = existing.get('Override_Strategy', '')
        enabled = existing.get('Enabled', 'FALSE')
        comment = existing.get('Comment', '')
        
        ws_rules.cell(row=row_idx, column=1, value=code).border = thin_border
        ws_rules.cell(row=row_idx, column=2, value=info.get('kpi_name', '')).border = thin_border
        ws_rules.cell(row=row_idx, column=3, value=info.get('kpi_type', '')).border = thin_border
        ws_rules.cell(row=row_idx, column=4, value=info.get('detection_strategy', '')).border = thin_border
        
        # Override_Strategy - editable (yellow)
        cell = ws_rules.cell(row=row_idx, column=5, value=override)
        cell.fill = edit_fill
        cell.border = thin_border
        dv_strategy.add(cell)
        
        # Enabled - editable (yellow)
        cell = ws_rules.cell(row=row_idx, column=6, value=enabled)
        cell.fill = edit_fill
        cell.border = thin_border
        dv_enabled.add(cell)
        
        # Change_Type
        cell = ws_rules.cell(row=row_idx, column=7, value=change_type)
        cell.border = thin_border
        if change_type == "Other→Regression":
            cell.fill = degraded_fill
        
        ws_rules.cell(row=row_idx, column=8, value=comment).border = thin_border
        
        row_idx += 1
    
    # Set column widths
    col_widths = [18, 55, 18, 18, 18, 10, 18, 30]
    for i, w in enumerate(col_widths, 1):
        ws_rules.column_dimensions[chr(64+i)].width = w
    
    ws_rules.freeze_panes = 'A2'
    ws_rules.auto_filter.ref = f"A1:H{row_idx-1}"
    
    # Sheet 2: Strategy Reference
    ws_ref = wb.create_sheet(title="2_Strategy_Reference")
    ws_ref['A1'] = "检测策略说明"
    ws_ref['A1'].font = Font(bold=True, size=14)
    
    ref_headers = ['Strategy', 'Description', 'Suitable_For']
    for col, h in enumerate(ref_headers, 1):
        cell = ws_ref.cell(row=3, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    ref_data = [
        ("STANDARD_SPC", "标准SPC控制图方法", "通用，默认方法"),
        ("ABS_THRESHOLD", "绝对阈值判定", "ceiling_saturated (success ratio=100%), floor_saturated (=0)"),
        ("RANGE_CONTEXT", "区间上下文判定 [μ±2σ]", "count_type (计数型KPI)"),
        ("RELATIVE_CHANGE", "相对变化率判定", "ratio_type (比例型，非饱和)"),
        ("ABS_DELTA_TREND", "绝对变化+Level Shift检测", "physical_stable (dB/dBm，低方差)"),
        ("QUANTILES", "分位数判定 (P50/P95)", "throughput_type (吞吐量)"),
        ("TAIL_AWARE", "尾部感知判定 (P95/P99)", "latency_type (延迟)"),
        ("WINDOW_COUNT", "滑窗计数判定", "low_frequency (低频事件)"),
        ("SEMANTIC_RULE", "语义规则判定", "config_type (配置/状态型)"),
    ]
    
    for r_idx, (strategy, desc, suitable) in enumerate(ref_data, 4):
        ws_ref.cell(row=r_idx, column=1, value=strategy).border = thin_border
        ws_ref.cell(row=r_idx, column=2, value=desc).border = thin_border
        ws_ref.cell(row=r_idx, column=3, value=suitable).border = thin_border
    
    ws_ref.column_dimensions['A'].width = 20
    ws_ref.column_dimensions['B'].width = 40
    ws_ref.column_dimensions['C'].width = 60
    
    # Sheet 3: Modification History
    ws_history = wb.create_sheet(title="3_Modification_History")
    
    hist_headers = ['Timestamp', 'Code', 'KPI_Name', 'Original_Strategy', 'New_Strategy', 'Modified_By']
    for col, h in enumerate(hist_headers, 1):
        cell = ws_history.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    # Load existing history
    if history_json.exists():
        try:
            with open(history_json, 'r', encoding='utf-8') as f:
                history = json.load(f)
            for h_idx, record in enumerate(history[-100:], 2):  # Keep last 100 records
                ws_history.cell(row=h_idx, column=1, value=record.get('timestamp', '')).border = thin_border
                ws_history.cell(row=h_idx, column=2, value=record.get('code', '')).border = thin_border
                ws_history.cell(row=h_idx, column=3, value=record.get('kpi_name', '')).border = thin_border
                ws_history.cell(row=h_idx, column=4, value=record.get('original_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=5, value=record.get('new_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=6, value=record.get('modified_by', 'user')).border = thin_border
        except:
            pass
    
    hist_widths = [20, 18, 50, 18, 18, 15]
    for i, w in enumerate(hist_widths, 1):
        ws_history.column_dimensions[chr(64+i)].width = w
    
    ws_history.freeze_panes = 'A2'
    
    # Sheet 4: Instructions
    ws_inst = wb.create_sheet(title="4_使用说明")
    instructions = [
        "KPI 自定义检测规则 - 使用说明",
        "",
        "【如何修改检测规则】",
        "1. 在 Sheet '1_Custom_Rules' 中找到要修改的 KPI (可用Filter筛选Change_Type)",
        "2. 在 'Override_Strategy' 列点击单元格，从下拉列表选择新策略",
        "3. 在 'Enabled' 列设置为 TRUE 启用该规则",
        "4. (可选) 在 'Comment' 列添加修改说明",
        "5. 保存 Excel 文件",
        "6. 运行命令: python scripts/manage_custom_rules.py --apply",
        "",
        "【列说明】",
        "- Code: KPI代码 (只读)",
        "- KPI_Name: KPI名称 (只读)",
        "- KPI_Type: 系统分类的KPI类型 (只读)",
        "- Default_Strategy: 系统默认的检测策略 (只读)",
        "- Override_Strategy: 【可编辑】用户自定义的检测策略",
        "- Enabled: 【可编辑】TRUE=启用自定义规则, FALSE=使用默认",
        "- Change_Type: 与旧版本对比的变化类型 (只读)",
        "- Comment: 【可编辑】修改说明",
        "",
        "【黄色背景】表示该列可编辑",
        "【红色背景】表示 Other→Regression (需要关注)",
        "",
        "【修改历史】",
        "- 每次运行 --apply 时，系统会自动记录修改历史",
        "- 历史记录保存在 Sheet '3_Modification_History'",
        "- 包含: 修改时间、原始策略、新策略",
    ]
    
    for r_idx, text in enumerate(instructions, 1):
        cell = ws_inst.cell(row=r_idx, column=1, value=text)
        if r_idx == 1:
            cell.font = Font(bold=True, size=14)
        elif text.startswith("【"):
            cell.font = Font(bold=True)
    
    ws_inst.column_dimensions['A'].width = 80
    
    wb.save(excel_path)
    
    print(f"\n自定义规则Excel已生成: {excel_path}")
    print(f"  - Sheet 1: {len(all_codes)} 个KPI的规则配置")
    print(f"  - Sheet 2: 检测策略说明")
    print(f"  - Sheet 3: 修改历史记录")
    print(f"  - Sheet 4: 使用说明")
    print(f"\n使用方法:")
    print(f"  1. 打开Excel，修改 Override_Strategy 和 Enabled 列")
    print(f"  2. 保存后运行: python scripts/manage_custom_rules.py --apply")


if __name__ == "__main__":
    main()
