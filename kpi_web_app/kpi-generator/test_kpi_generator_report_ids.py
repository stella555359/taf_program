import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook

from kpi_generator import (
    CompassClient,
    FinalReportRequestError,
    KpiGeneratorRequest,
    KpiGeneratorService,
    TimeRange,
    build_interval_package_version,
)


class CompassClientReportIdFormattingTests(unittest.TestCase):
    def test_compare_report_ids_are_not_wrapped_in_quotes(self) -> None:
        self.assertEqual(
            CompassClient._format_compare_report_ids(['1534183306', "'1534183307'", ' 1534183308 ']),
            '1534183306&1534183307&1534183308',
        )

    def test_horizontal_merge_report_id_list_is_numeric_literal_text(self) -> None:
        self.assertEqual(
            CompassClient._format_report_id_list_literal(['1534183306', '"1534183307"']),
            '["1534183306", "1534183307"]',
        )

    def test_final_report_json_error_raises_with_request_data(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {'Content-Type': 'application/json'}
        response.status_code = 200
        response.url = client.COMPARE_URL
        response.json.return_value = {
            'success': False,
            'error': 'Field \'id\' expected a number but got "\'1534183306\'".',
            'data': None,
        }

        with TemporaryDirectory() as temp_dir:
            with self.assertRaises(FinalReportRequestError) as context:
                client._save_response_to_path(
                    response,
                    Path(temp_dir) / 'report.xlsx',
                    fallback_sheet_name='Compare Report',
                    operation='compare',
                    report_ids=['1534183306'],
                    request_data={'reports': '1534183306'},
                )

        self.assertEqual(context.exception.diagnostics['request_data'], {'reports': '1534183306'})
        self.assertIn('expected a number', context.exception.diagnostics['response_body'])

    def test_compare_report_json_is_converted_to_detector_compatible_workbook(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {'Content-Type': 'application/json'}
        response.status_code = 200
        response.url = client.COMPARE_URL
        response.json.return_value = {
            'packages': ['T813.20260401_010000', 'T813.20260401_020000', 'T813.20260401_030000'],
            'data': [
                {
                    'group_name': 'DL',
                    'code': 'KPI-001',
                    'kpi_name': 'Throughput',
                    'value1': 11,
                    'value2': 22,
                    'value3': 33,
                }
            ],
        }

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'compare.xlsx'
            saved_path = client._save_response_to_path(
                response,
                output_path,
                fallback_sheet_name='Compare Report',
                operation='compare',
                report_ids=['1534183306', '1534183307', '1534183308'],
                request_data={'reports': '1534183306&1534183307&1534183308'},
                environment='T813.SCF.T813.demo',
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            workbook = load_workbook(saved_path)
            sheet = workbook['KPI Report']

            self.assertEqual(saved_path, output_path)
            self.assertEqual(
                [sheet.cell(row=1, column=index).value for index in range(1, 7)],
                [
                    'Group Name',
                    'Code',
                    'KPI Name',
                    'T813.20260401_010000',
                    'T813.20260401_020000',
                    'T813.20260401_030000',
                ],
            )
            self.assertEqual([sheet['A2'].value, sheet['B2'].value, sheet['C2'].value], ['DL', 'KPI-001', 'Throughput'])
            self.assertEqual([sheet['D2'].value, sheet['E2'].value, sheet['F2'].value], [11, 22, 33])
            self.assertEqual(sheet.column_dimensions['A'].width, 34)
            self.assertEqual(sheet.column_dimensions['B'].width, 16)
            self.assertEqual(sheet.column_dimensions['C'].width, 53)
            self.assertEqual(sheet.column_dimensions['D'].width, 18)
            self.assertEqual(sheet.row_dimensions[1].height, 83)
            self.assertEqual(sheet.row_dimensions[2].height, 31)

    def test_compare_report_non_detector_package_names_are_forced_to_detector_style(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {'Content-Type': 'application/json'}
        response.status_code = 200
        response.url = client.COMPARE_URL
        response.json.return_value = {
            'packages': ['first package', 'second package'],
            'data': [
                {
                    'group_name': 'DL',
                    'code': 'KPI-001',
                    'kpi_name': 'Throughput',
                    'value1': 11,
                    'value2': 22,
                }
            ],
        }

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'compare_forced.xlsx'
            saved_path = client._save_response_to_path(
                response,
                output_path,
                fallback_sheet_name='Compare Report',
                operation='compare',
                report_ids=['1534183306', '1534183307'],
                request_data={'reports': '1534183306&1534183307'},
                environment='T813.SCF.T813.demo',
                interval_timestamps=[datetime(2026, 4, 1, 1, 0, 0), datetime(2026, 4, 1, 2, 0, 0)],
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            workbook = load_workbook(saved_path)
            sheet = workbook['KPI Report']

            self.assertEqual(
                [sheet.cell(row=1, column=index).value for index in range(1, 6)],
                ['Group Name', 'Code', 'KPI Name', 'T813.20260401_010000', 'T813.20260401_020000'],
            )
            self.assertEqual([sheet['D2'].value, sheet['E2'].value], [11, 22])

    def test_compare_report_preserves_existing_detector_style_headers(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {'Content-Type': 'application/json'}
        response.status_code = 200
        response.url = client.COMPARE_URL
        response.json.return_value = {
            'packages': ['T813.20260401_010000.26R2.0126.7', 'T813.20260401_020000.26R2.0126.7'],
            'data': [
                {
                    'group_name': 'DL',
                    'code': 'KPI-001',
                    'kpi_name': 'Throughput',
                    'value1': 11,
                    'value2': 22,
                }
            ],
        }

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'compare_preserved.xlsx'
            saved_path = client._save_response_to_path(
                response,
                output_path,
                fallback_sheet_name='Compare Report',
                operation='compare',
                report_ids=['1534183306', '1534183307'],
                request_data={'reports': '1534183306&1534183307'},
                environment='T813.SCF.T813.demo',
                interval_timestamps=[datetime(2026, 4, 1, 1, 0, 0), datetime(2026, 4, 1, 2, 0, 0)],
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            workbook = load_workbook(saved_path)
            sheet = workbook['KPI Report']

            self.assertEqual(
                [sheet.cell(row=1, column=index).value for index in range(1, 6)],
                ['Group Name', 'Code', 'KPI Name', 'T813.20260401_010000.26R2.0126.7', 'T813.20260401_020000.26R2.0126.7'],
            )

    def test_analysis_report_uses_detector_style_single_data_column(self) -> None:
        client = CompassClient()
        session = Mock()
        response = Mock()
        response.json.return_value = {
            'data': [
                {
                    'group_name': 'DL',
                    'code': 'KPI-001',
                    'kpi_name': 'Throughput',
                    'value': 123.45,
                }
            ]
        }
        session.post.return_value = response
        client.session = session

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'analysis.xlsx'
            saved_path = client.get_analysis_kpi_result(
                '1534183306',
                output_path,
                environment='T813.SCF.T813.demo',
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            workbook = load_workbook(saved_path)
            sheet = workbook['KPI Report']

            self.assertEqual(saved_path, output_path)
            self.assertEqual(
                [sheet.cell(row=1, column=index).value for index in range(1, 5)],
                ['Group Name', 'Code', 'KPI Name', 'T813.20260407_123456'],
            )
            self.assertEqual([sheet['A2'].value, sheet['B2'].value, sheet['C2'].value], ['DL', 'KPI-001', 'Throughput'])
            self.assertEqual(sheet['D2'].value, 123.45)
            self.assertEqual(sheet.column_dimensions['A'].width, 34)
            self.assertEqual(sheet.column_dimensions['B'].width, 16)
            self.assertEqual(sheet.column_dimensions['C'].width, 53)
            self.assertEqual(sheet.column_dimensions['D'].width, 18)

    def test_horizontal_merge_attachment_is_normalized_to_detector_columns(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {
            'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        }
        response.status_code = 200
        response.url = client.HORIZONTAL_MERGE_URL

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sheet1'
        sheet.append([
            'group_name',
            'code',
            'kpi_name',
            'SBTS26R3.ENB.9999.260319.000005',
            'SBTS26R3.ENB.9999.260319.000005',
            'SBTS26R3.ENB.9999.260319.000005',
            'SBTS26R3.ENB.9999.260319.000005',
        ])
        sheet.append(['DL', 'KPI-001', 'Throughput', 11, 22, 33, 44])
        buffer = BytesIO()
        workbook.save(buffer)
        response.content = buffer.getvalue()

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'horizontal_merge.xlsx'
            saved_path = client._save_response_to_path(
                response,
                output_path,
                fallback_sheet_name='Horizontal Merge',
                operation='horizontal_merge',
                report_ids=['1534183306', '1534183307', '1534183308', '1534183309'],
                request_data={'report_id_list': '["1534183306", "1534183307", "1534183308", "1534183309"]'},
                environment='T813.SCF.T813.demo',
                interval_timestamps=[
                    datetime(2026, 4, 1, 1, 0, 0),
                    datetime(2026, 4, 1, 2, 0, 0),
                    datetime(2026, 4, 1, 3, 0, 0),
                    datetime(2026, 4, 1, 4, 0, 0),
                ],
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            normalized_workbook = load_workbook(saved_path)
            normalized_sheet = normalized_workbook['KPI Report']

            self.assertEqual(saved_path, output_path)
            self.assertEqual(
                [normalized_sheet.cell(row=1, column=index).value for index in range(1, 8)],
                [
                    'Group Name',
                    'Code',
                    'KPI Name',
                    'T813.20260401_010000',
                    'T813.20260401_020000',
                    'T813.20260401_030000',
                    'T813.20260401_040000',
                ],
            )
            self.assertEqual(
                [normalized_sheet['D2'].value, normalized_sheet['E2'].value, normalized_sheet['F2'].value, normalized_sheet['G2'].value],
                [11, 22, 33, 44],
            )
            self.assertEqual(normalized_sheet.column_dimensions['D'].width, 18)
            self.assertEqual(normalized_sheet.row_dimensions[1].height, 83)

    def test_horizontal_merge_preserves_existing_detector_style_headers(self) -> None:
        client = CompassClient()
        response = Mock()
        response.headers = {
            'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        }
        response.status_code = 200
        response.url = client.HORIZONTAL_MERGE_URL

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sheet1'
        sheet.append([
            'group_name',
            'code',
            'kpi_name',
            'T813.20260401_010000.26R2.0126.7',
            'T813.20260401_020000.26R2.0126.7',
        ])
        sheet.append(['DL', 'KPI-001', 'Throughput', 11, 22])
        buffer = BytesIO()
        workbook.save(buffer)
        response.content = buffer.getvalue()

        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'horizontal_merge_preserved.xlsx'
            saved_path = client._save_response_to_path(
                response,
                output_path,
                fallback_sheet_name='Horizontal Merge',
                operation='horizontal_merge',
                report_ids=['1534183306', '1534183307'],
                request_data={'report_id_list': '["1534183306", "1534183307"]'},
                environment='T813.SCF.T813.demo',
                interval_timestamps=[
                    datetime(2026, 4, 1, 1, 0, 0),
                    datetime(2026, 4, 1, 2, 0, 0),
                ],
                generated_at=datetime(2026, 4, 7, 12, 34, 56),
            )

            normalized_workbook = load_workbook(saved_path)
            normalized_sheet = normalized_workbook['KPI Report']

            self.assertEqual(
                [normalized_sheet.cell(row=1, column=index).value for index in range(1, 6)],
                [
                    'Group Name',
                    'Code',
                    'KPI Name',
                    'T813.20260401_010000.26R2.0126.7',
                    'T813.20260401_020000.26R2.0126.7',
                ],
            )


class KpiGeneratorIntervalNamingTests(unittest.TestCase):
    def test_build_interval_package_version_includes_detector_prefix(self) -> None:
        request = KpiGeneratorRequest(
            template_set_name='setA',
            template_names=['templateA'],
            build='26R2.0126.7',
            environment='T813.SCF.T813.demo',
            scenario='burst',
            report_timestamps_list=[TimeRange(start=datetime(2026, 4, 1, 1, 0, 0), end=datetime(2026, 4, 1, 2, 0, 0))],
            test_line='7_5_UTE5G402T813',
        )

        package_version = build_interval_package_version(
            request,
            TimeRange(start=datetime(2026, 4, 1, 1, 0, 0), end=datetime(2026, 4, 1, 2, 0, 0)),
        )

        self.assertEqual(package_version, 'T813.20260401_010000.26R2.0126.7')

    def test_generate_interval_bundle_uses_same_package_version_within_interval(self) -> None:
        client = Mock()
        client.generate_new_report.side_effect = ['101', '102']
        client.combine_kpi_report.return_value = '201'
        service = KpiGeneratorService(client=client, output_dir=Path('.'), logger=Mock())
        interval = TimeRange(start=datetime(2026, 4, 1, 1, 0, 0), end=datetime(2026, 4, 1, 2, 0, 0))
        request = KpiGeneratorRequest(
            template_set_name='setA',
            template_names=['templateA', 'templateB'],
            build='26R2.0126.7',
            environment='T813.SCF.T813.demo',
            scenario='burst',
            report_timestamps_list=[interval],
            test_line='7_5_UTE5G402T813',
        )

        combined_id, report_ids, failed_templates, interval_detail = service._generate_interval_bundle(
            request,
            interval,
            ['templateA', 'templateB'],
            1,
        )

        self.assertEqual(combined_id, '201')
        self.assertEqual(report_ids, ['101', '102'])
        self.assertEqual(failed_templates, [])
        self.assertEqual(interval_detail['status'], 'completed')
        self.assertEqual(client.generate_new_report.call_count, 2)
        self.assertEqual(client.generate_new_report.call_args_list[0].kwargs['package_version'], 'T813.20260401_010000.26R2.0126.7')
        self.assertEqual(client.generate_new_report.call_args_list[1].kwargs['package_version'], 'T813.20260401_010000.26R2.0126.7')


if __name__ == '__main__':
    unittest.main()