from openpyxl import Workbook


def create_excel_workbook(extracted_data, output_file):
    workbook = Workbook()

    # Remove the default sheet created by Workbook
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheet = workbook.create_sheet(title="ExtractedData")
    for i, value in enumerate(extracted_data, start=1):
        sheet.cell(row=i, column=1, value=value)

    workbook.save(output_file)


def extract_data(file_path, keyword):
    extracted_data = []
    with open(file_path, 'r') as file:
        for line in file:
            if keyword in line:
                extracted_data.append(line.strip())
    return extracted_data


def process_file(input_file, keyword, output_file):
    extracted_data = extract_data(input_file, keyword)
    create_excel_workbook(extracted_data, output_file)
    return f"Extracted data has been saved to {output_file}."
