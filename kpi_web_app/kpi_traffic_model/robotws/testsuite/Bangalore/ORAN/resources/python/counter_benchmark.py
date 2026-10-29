import openpyxl
from openpyxl import load_workbook
import json
import os
import datetime
from datetime import timedelta, datetime
from xml.etree import ElementTree as ET
from taf.sbts.oam.coam.admin import admin
import time
from taf.analyzer.pmcounter import PMCounterLogReader
import sys


class counter_benchmark():
    def __init__(self):
        pass

    def get_admin_pminterval(self, bts):
        api = admin()
        conn = api.setup_admin(bts_host=bts['host'], username=bts['user'], password=bts['password'])
        api.get_current_scf_file("/home/ute/counters", "scf1.xml", connection=conn)
        tree = ET.parse("/home/ute/counters/scf1.xml")
        root = tree.getroot()[0]
        class_name = "reportingIntervalPm"
        interval = ""
        child_lischild_list = root.findall(".//*[@name=\"" + class_name + "\"]")
        for ch in child_lischild_list:
            interval_text = ch.text
            interval = interval_text.replace("min", "")
        sw_version = api.get_sw_version()
        api.teardown(connection=conn)
        return interval, sw_version

    def wait_for_the_log(self, starttime, endtime, interval):
        format = '%Y-%m-%d %H:%M:%S'
        interval = int(interval)
        start_time1 = starttime.split(".")
        starttime = datetime.strptime(start_time1[0], format)
        end_time1 = endtime.split(".")
        endtime = datetime.strptime(end_time1[0], format)
        modulus = int(starttime.minute) % interval
        print(modulus)
        if modulus > 0:
            start_future_date = starttime + timedelta(minutes=(interval - modulus))
        else:
            start_future_date = starttime
        print(start_future_date)

        modulus = int(endtime.minute) % interval
        print(modulus)
        if modulus > 0:
            end_future_date = endtime + timedelta(minutes=(interval - modulus))
        else:
            end_future_date = endtime
        print(end_future_date)

        if end_future_date < datetime.now():
            print("%%%%%%")
        else:
            s = end_future_date - datetime.now()
            wait_sec = int(s.seconds) + 20
            time.sleep(wait_sec)

        for_last = end_future_date - starttime
        for_last_sec = for_last.seconds
        return for_last_sec

    def write_counter(self, build, test_case, log_entry, counter_file, start_time, end_time):
        f = open(counter_file)
        data = json.load(f)
        f.close()
        result = None

        validation_err = []
        print(test_case, build, data, log_entry)

        for cell in data.keys():
            dn_regex = "/" + cell
            for counter_detail in data[cell]:
                if counter_detail["counter"]:
                    ckey = counter_detail["counter"]
                    value = ""
                    c_value = 0
                    print("#######", log_entry)
                    for logentry in log_entry:
                        try:
                            c_value += logentry.get_pi_value(pi_id=ckey, dn_re=dn_regex)
                            print("&&&&&&&&&&&&&&&&&&&&")
                            print(ckey)
                            print(c_value)
                            value = c_value
                        except Exception as e:
                            print(e)
                            continue
                    time_stamp = str(start_time) + "-" + str(end_time)
                    print(value, time_stamp)
                    if value != "":
                        result = self._write_to_xl(build, test_case, counter_detail, time_stamp, value, cell)
                    if result is not None:
                        validation_err.append(result)
        return validation_err

    def _write_to_xl(self, build, tc_name, counter_detail, timestamp, value, cell_val):
        result = "Valid"
        reason = "-"
        print("*****************", build, tc_name, counter_detail, timestamp, value, cell_val)
        counter_name = counter_detail['counter']
        timestamp = str(timestamp)
        if not os.path.isfile("/home/ute/counters/counter.xlsx"):
            wb = openpyxl.Workbook()
            wb.save("/home/ute/counters/counter.xlsx")
        workbook = load_workbook("/home/ute/counters/counter.xlsx")
        if "counters" not in workbook.sheetnames:
            workbook.create_sheet('counters')
            sheet = workbook['counters']
            sheet.cell(row=1, column=1).value = "tc_name"
            sheet.cell(row=1, column=2).value = "Cell"
            workbook.save("/home/ute/counters/counter.xlsx")
        if tc_name not in workbook.sheetnames:
            workbook.create_sheet(tc_name)
            sheet = workbook[tc_name]
            sheet.cell(row=1, column=1).value = "Build"
            sheet.cell(row=1, column=2).value = "Timestamp"
            sheet.cell(row=1, column=3).value = "Result"
            sheet.cell(row=1, column=4).value = "Reason"
            sheet.cell(row=1, column=5).value = "Cell"

            workbook.save("/home/ute/counters/counter.xlsx")
        sheet = workbook[tc_name]
        if counter_name not in [cell.value for cell in sheet[1]]:
            sheet.cell(1, sheet.max_column + 1).value = counter_name
            workbook.save("/home/ute/counters/counter.xlsx")
        valid_rows = [row for row in sheet.iter_rows(values_only=True) if row[4] == cell_val and row[2] == "Valid"]
        column_index = None
        for i, cell in enumerate(sheet[1]):
            if cell.value == counter_name:
                column_index = i
                break
        if column_index is not None:
            if len([row for row in valid_rows if row[column_index] is not None]) > 1:
                self._write_min_max_values(valid_rows, counter_name, tc_name, cell_val)
                workbook = load_workbook("/home/ute/counters/counter.xlsx")
        sheet_counters = workbook["counters"]
        min_max_col = None
        min_max_rows = [row for row in sheet_counters.iter_rows(values_only=True)]
        for i, cell in enumerate(sheet_counters[1]):
            if cell.value == counter_name:
                min_max_col = i
                break

        if counter_detail["max_val"] == "" or counter_detail["min_val"] == "":
            if column_index is not None:
                if len([row for row in valid_rows if row[column_index] is not None]) > 1:
                    max_val = max(row[column_index] for row in valid_rows if row[column_index] is not None)
                    min_val = min(row[column_index] for row in valid_rows if row[column_index] is not None)
                else:
                    max_val = ""
                    min_val = ""
        else:
            max_val = counter_detail["max_val"]
            min_val = counter_detail["min_val"]
        if counter_detail["deviation"] == "":
            deviation = 0.15
        else:
            deviation = counter_detail["deviation"] / 100

        if max_val != "" and min_val != "":
            if len([row for row in valid_rows if row[column_index] is not None]) > 1:
                max_limit = max_val + (max_val * deviation)
                min_limit = min_val - (min_val * deviation)
                if value < min_limit or value > max_limit:
                    reason = reason + " found counter " + str(counter_name) + " deviation more than " + str(
                        deviation) + "."
                    result = "Not Valid"
                else:
                    minmax_df = [row for row in min_max_rows if row[0] == tc_name and row[1] == cell_val]
                    initial_min_max = minmax_df[0][min_max_col].split(",")
                    if (int(float(value)) > int(float(max_val))) and (
                            int(float(value)) > (
                            int(float(initial_min_max[0])) + int(float(initial_min_max[0])) * 0.3)):
                        reason = reason + str(counter_name) + " value: " + str(
                            value) + " deviate more than 30% from initial max value " + initial_min_max[0] + "."
                        result = "Not Valid"
                    if (int(float(value)) < int(float(min_val))) and (
                            int(float(value)) < (
                            int(float(initial_min_max[1])) - int(float(initial_min_max[1])) * 0.3)):
                        reason = reason + str(counter_name) + " value: " + str(
                            value) + " deviate more than 30% from initial min value " + initial_min_max[0] + "."
                        result = "Not Valid"

        sheet = workbook[tc_name]
        new_df_filter = [row for row in sheet.iter_rows(values_only=True) if
                         row[4] == cell_val and row[1] == timestamp and row[0] == build]
        if len(new_df_filter) > 0:
            existing_reason = new_df_filter[0][3]
            if reason != "-":
                final_reason = str(existing_reason) + str(reason).replace("-", "")
            else:
                final_reason = str(existing_reason)

            row_index = None
            for j, row in enumerate(sheet.iter_rows(values_only=True)):
                if row[4] == cell_val and row[1] == timestamp and row[0] == build:
                    row_index = j + 1
                    break

            sheet.cell(row_index, column_index + 1).value = value
            sheet.cell(row_index, 4).value = final_reason
            if sheet.cell(row_index, 3).value == "Valid":
                sheet.cell(row_index, 3).value = result

        else:
            column_index = None
            for i, cell in enumerate(sheet[1]):
                if cell.value == counter_name:
                    column_index = i + 1
                    break
            if column_index is not None:
                new_row_index = sheet.max_row + 1
                sheet.cell(row=new_row_index, column=column_index).value = value
            # Save the workbook
            sheet.cell(new_row_index, 1).value = build
            sheet.cell(new_row_index, 2).value = timestamp
            sheet.cell(new_row_index, 5).value = cell_val
            sheet.cell(new_row_index, 3).value = result
            sheet.cell(new_row_index, 4).value = reason
        workbook.save("/home/ute/counters/counter.xlsx")

    def _write_min_max_values(self, valid_rows, counter_name, tc_name, cell_val):
        workbook = load_workbook("/home/ute/counters/counter.xlsx")
        sheet_counters = workbook['counters']
        min_max_rows = [row for row in sheet_counters.iter_rows(values_only=True)]
        min_max_row_len = len([row for row in min_max_rows if row[0] == tc_name and row[1] == cell_val])
        if min_max_row_len == 0:
            row = sheet_counters.max_row + 1
            sheet_counters.cell(row, 1).value = tc_name
            sheet_counters.cell(row, 2).value = cell_val

        if counter_name not in [cell.value for cell in sheet_counters[1]]:
            sheet_counters.cell(1, sheet_counters.max_column + 1).value = counter_name
        workbook.save("/home/ute/counters/counter.xlsx")
        min_max_rows = [row for row in sheet_counters.iter_rows(values_only=True) if
                        row[0] == tc_name and row[1] == cell_val]
        column_index = None
        for i, cell in enumerate(sheet_counters[1]):
            if cell.value == counter_name:
                column_index = i
                break

        sheet = workbook[tc_name]
        column_index1 = None
        for i, cell in enumerate(sheet[1]):
            if cell.value == counter_name:
                column_index1 = i
                break

        row_index = None
        for j, row in enumerate(sheet_counters.iter_rows(values_only=True)):
            if row[0] == tc_name and row[1] == cell_val:
                row_index = j + 1
                break

        if min_max_rows[0][column_index] is None:
            max_val = max(row[column_index1] for row in valid_rows if row[column_index1] is not None)
            min_val = min(row[column_index1] for row in valid_rows if row[column_index1] is not None)
            val = str(max_val) + "," + str(min_val)
            sheet_counters.cell(row_index, column_index + 1).value = val
            workbook.save("/home/ute/counters/counter.xlsx")
        workbook.save("/home/ute/counters/counter.xlsx")
