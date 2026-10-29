"""
:modified:05.01.2023

:author: Mateusz Kawka
:maintainer: Mateusz Kawka
:contact: mateusz.kawka@nokia.com
"""

import warnings
import csv
import os
import re
# import pandas as pd
from robot.api.deco import keyword
from robot.api import logger
from openpyxl.workbook import Workbook
from openpyxl.styles import Color, PatternFill, Font, Border, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.comments import Comment
from openpyxl import load_workbook

warnings.filterwarnings(action="ignore", message="unclosed")
warnings.filterwarnings(action="ignore", message="ImportWarning:")

ROBOT_AUTO_KEYWORDS = False


@keyword(name='Create Excel File')
def create_excel_file(file_name, path_to_file, sheet_name=None, headers=None):
    if not os.path.isdir(path_to_file):
        os.makedirs(path_to_file)
    workbook_name = file_name + '.xlsx'
    wb = Workbook()
    sheet = wb.active
    if sheet_name:
        sheet.title = sheet_name
    if headers:
        headers = headers.split(',')
        sheet.append(headers)
    wb.save(path_to_file + '/' + workbook_name)
    print("XLSX file {0} created in {1}".format(file_name, path_to_file))


@keyword(name='Append Next Row To Excel File')
def append_next_raw_to_excel_file(path_to_file, list_of_data, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    sheet.append(list_of_data)
    print("Data: {0} appended to {1}".format(list_of_data, sheet_name))
    wb.save(path_to_file)


@keyword(name='Append Column With Data To Excel File')
def append_column_with_data_to_excel_file(path_to_file, list_of_data, start_cell, header=None, sheet_name=None):
    column = re.findall(r'[a-zA-Z]+', start_cell)[0]
    row = int(re.findall(r'\d+', start_cell)[0])
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    if header:
        sheet[start_cell] = header
        row += 1
    for i in list_of_data:
        sheet[column + str(row)] = i
        print('Cell {0} with value {1}'.format(column + str(row), i))
        row += 1
    wb.save(path_to_file)


@keyword(name='Append To Specified Cell Excel File')
def append_to_specified_cell_excel_file(path_to_file, cell_value,
                                        start_row, end_row, start_column, end_column, bold=False, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    if start_row != end_row or start_column != end_column:
        sheet.merge_cells(start_row=start_row, start_column=start_column, end_row=end_row, end_column=end_column)
    sheet.cell(row=start_row, column=start_column).value = cell_value
    sheet.cell(row=start_row, column=start_column).alignment = Alignment(horizontal='center', vertical='center')
    sheet.cell(row=start_row, column=start_column).font = Font(bold=bold)
    wb.save(path_to_file)


@keyword(name='Get First Empty Column Excel File')
def get_first_empty_column_excel_file(path_to_file, int_or_letter, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    if sheet_name is None:
        sheet = wb.active
        if int_or_letter == 'int':
            return sheet.max_column + 1
        else:
            return get_column_letter(sheet.max_column + 1)
    elif sheet_name not in wb.sheetnames:
        print("No sheet '{0}' in excel file !".format(sheet_name))
        return None
    else:
        sheet = wb[sheet_name]
        print("Given sheet '{0}' exists".format(sheet_name))
        if int_or_letter == 'int':
            return sheet.max_column + 1
        else:
            return get_column_letter(sheet.max_column + 1)


@keyword(name="Get Last Row For Column")
def get_last_row_for_column(path_to_file, column, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    max_row = len(sheet[column])
    return max_row


@keyword(name='Add Comment To Cell')
def add_comment_to_cell(path_to_file, cell, comment, author='TA', sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    comment = Comment(comment, author)
    sheet[cell].comment = comment
    wb.save(path_to_file)


@keyword(name='Get Columns Values')
def get_columns_values(path_to_file, min_row, max_row, min_col, max_col, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    cols_iter = sheet.iter_cols(min_col=int(min_col), min_row=int(min_row), max_col=int(max_col), max_row=int(max_row))
    values_lists = [[cell.value for cell in col] for col in cols_iter]

    return values_lists


@keyword(name='Find Cell By Value')
def find_cell_by_value(path_to_file, value, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    for row in sheet.iter_rows(min_col=1, min_row=1, max_col=sheet.max_column, max_row=sheet.max_row):
        for cell in row:
            if cell.value == value:
                return get_column_letter(cell.column) + str(cell.row)


@keyword(name='Get Cell Value')
def get_cell_value(path_to_file, cell, sheet_name=None):
    path_to_file = path_to_file + '.xlsx'
    wb = load_workbook(path_to_file)
    sheet = _get_sheet(wb, sheet_name)
    return sheet[cell].value


def _get_sheet(workbook, sheet_name):
    if sheet_name is None:
        sheet = workbook.active
        print("Sheet name is not given, active is used")
    elif sheet_name not in workbook.sheetnames:
        workbook.create_sheet(sheet_name)
        sheet = workbook[sheet_name]
        print("Created new sheet {0}".format(sheet_name))
    else:
        sheet = workbook[sheet_name]
        print("Given sheet ${0} exists".format(sheet_name))

    return sheet


# @keyword(name='Create CSV file')
# def create_csv_file(path_to_file):
#     with open(path_to_file, 'w') as csvfile:
#         filewriter = csv.writer(csvfile, quoting=csv.QUOTE_MINIMAL)
#         filewriter.writerow(['Variant', 'KPI value'])
#         print("Created new CSV file in path {0}".format(path_to_file))
#         csvfile.close()


# @keyword(name='Append To CSV file')
# def append_to_csv_file(path_to_file, name, value):
#     with open(path_to_file, 'a') as csvfile:
#         filewriter = csv.writer(csvfile, quoting=csv.QUOTE_MINIMAL)
#         filewriter.writerow([name, value])
#         print("Variant: {0} saved with value: {1}".format(name, value))
#         csvfile.close()


# @keyword(name='Convert CSV To Excel File')
# def convert_csv_to_excel_file(csv_file_path, xlsx_file_path):
#     csv_to_df = pd.read_csv(csv_file_path)
#     writer = pd.ExcelWriter(xlsx_file_path)
#     csv_to_df.to_excel(writer, index=False)
#     writer.save()
#     logger.info("File {0} converted to {1} successfully !".format(csv_file_path, xlsx_file_path))
