import requests
import json
import re
import ast
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
from openpyxl.styles.borders import Border, Side

_required_params = ["feature", "not_automated"]
_required_params2 = ["name", "feature", "det_auto_lvl", "test_entity", "status", "res_tester", "release"]

thin_border = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))


def get_data(url):
    response = requests.get(url)
    data = response.json()
    return data


def create_xml(data, filename):
    wb = Workbook()
    dest_filename = filename
    ws1 = wb.active
    ws1.title = 'Feature_List'
    starting_cell_fill = 1
    for i in _required_params:
        ws1.cell(column=starting_cell_fill, row=1, value=str(i)).border = thin_border
        starting_cell_fill += 1
    starting_cell_fill = 1
    current_row = 2
    for tc_dict in data:
        for i in _required_params:
            print(tc_dict)
            ws1.cell(column=starting_cell_fill, row=current_row, value=tc_dict.get(i, None)).border = thin_border
            starting_cell_fill += 1
        current_row += 1
        starting_cell_fill = 1
    wb.save(filename=dest_filename)


def new_entry(data_log, filename2):
    wb2 = load_workbook(filename2)
    dest_filename = filename2
    ws2 = wb2.create_sheet("Test_list")
    starting_cell_fill = 1
    for i in _required_params2:
        ws2.cell(column=starting_cell_fill, row=1, value=str(i)).border = thin_border
        starting_cell_fill += 1
    starting_cell_fill = 1
    current_row = 2
    for tc_dict in data_log:
        for i in _required_params2:
            print(tc_dict)
            ws2.cell(column=starting_cell_fill, row=current_row, value=tc_dict.get(i, None)).border = thin_border
            starting_cell_fill += 1
        current_row += 1
        starting_cell_fill = 1
    wb2.save(filename=dest_filename)
