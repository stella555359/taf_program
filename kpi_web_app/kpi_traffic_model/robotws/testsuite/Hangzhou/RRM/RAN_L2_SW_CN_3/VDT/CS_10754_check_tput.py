#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import paramiko
import sys
import os
from datetime import datetime, timedelta
import tempfile

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    # ======================
    # Linux SSH 信息
    # ======================
    host = "10.57.251.84"
    port = 22
    username = "ute"
    password = "uteUTE!"

    # ======================
    # Windows SFTP 信息
    # ======================
    win_host = "10.57.130.195"
    win_port = 22
    win_user = r"noklab\SG-tester"
    win_pass = "Nokia123!"
    win_dir = "/D:/5g_log/crt/10754"

    # ======================
    # Linux 测试目录
    # ======================
    remote_dir = (
        "/home/ute/MN/UTE/robotws/testsuite/"
        "Hangzhou/RRM/RAN_L2_SW_CN_3/VDT"
    )

    pyenv_env = "test_repo"
    robot_cmd = "robot -L trace VDT_10754_check_gnb_tput.robot"
    ts = (datetime.now() - timedelta(hours=8)).strftime("%Y%m%d_%H%M%S")

    full_cmd = f"""
export PYENV_ROOT="$HOME/.pyenv"
export PATH="$PYENV_ROOT/bin:$PATH"
eval "$(pyenv init -)"
pyenv activate {pyenv_env}

cd {remote_dir}
# ===== 执行前清理旧 Snapshot 和 log 目录 =====
rm -f Snapshot_*.zip
rm -rf log
rm -rf admin_*
rm -rf output*
rm -rf report.html
rm 10754_kpi*

touch .snapshot_marker_{ts}
{robot_cmd}
rc=$?

if [ $rc -eq 0 ]; then
    result=pass
else
    result=fail
fi

if [ -f log.html ]; then
    mv log.html log_{ts}_${{result}}.html
fi

echo RESULT=$result
SNAPSHOT_FILES=$(find . -maxdepth 1 -type f -name 'Snapshot_*.zip' -newer .snapshot_marker_{ts})
echo $SNAPSHOT_FILES
"""

    try:
        # ======================
        # 连接 Linux
        # ======================
        print(f"[INFO] 连接 Linux: {host}")
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        ssh.connect(host, port, username, password)

        stdin, stdout, stderr = ssh.exec_command(full_cmd)
        stdout.channel.recv_exit_status()
        out = stdout.read().decode()
        err = stderr.read().decode()
        print(out)
        print(err)

        # 解析 pass / fail
        result = "pass"
        for line in out.splitlines():
            if line.startswith("RESULT="):
                result = line.split("=", 1)[1].strip()

        sftp = ssh.open_sftp()

        # ======================
        # 连接 Windows
        # ======================
        print("[INFO] 连接 Windows PC")
        win_ssh = paramiko.SSHClient()
        win_ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        win_ssh.connect(win_host, win_port, win_user, win_pass)
        win_sftp = win_ssh.open_sftp()

        # ======================
        # 上传 log 文件
        # ======================
        log_name = f"log_{ts}_{result}.html"
        with sftp.file(f"{remote_dir}/{log_name}", "rb") as f:
            win_sftp.putfo(f, f"{win_dir}/{log_name}")

        # ======================
        # 上传 Snapshot 文件
        # ======================
        for line in out.splitlines():
            if line.startswith("./Snapshot_") and line.endswith(".zip"):
                with sftp.file(f"{remote_dir}/{line[2:]}", "rb") as f:
                    win_sftp.putfo(f, f"{win_dir}/{line[2:]}")

        # ======================
        # KPI 文件重命名并上传
        # ======================
        kpi_src = f"{remote_dir}/10754_kpi.xls"
        kpi_new = f"10754_kpi_{ts}_{result}.xls"
        kpi_dst = f"{remote_dir}/{kpi_new}"

        try:
            sftp.stat(kpi_src)
            sftp.rename(kpi_src, kpi_dst)

            with sftp.file(kpi_dst, "rb") as f:
                win_sftp.putfo(f, f"{win_dir}/{kpi_new}")

            print("[INFO] KPI 上传完成")

        except IOError:
            print("[INFO] 无 KPI 文件，跳过上传")

        # ======================
        # KPI 汇总到 Total（新增 timestamp / result + 自动着色）
        # ======================
        print("[INFO] 开始汇总 KPI 到 10754_Total_KPI.xlsx")

        total_kpi_xlsx = f"{win_dir}/10754_Total_KPI.xlsx"

        pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        with tempfile.TemporaryDirectory() as tmpdir:
            local_total = os.path.join(tmpdir, "10754_Total_KPI.xlsx")
            local_kpi = os.path.join(tmpdir, kpi_new)

            # 下载本次 KPI
            win_sftp.get(f"{win_dir}/{kpi_new}", local_kpi)

            # 判断 Total 是否存在
            try:
                win_sftp.get(total_kpi_xlsx, local_total)
                wb_total = load_workbook(local_total)
                ws_total = wb_total.active
                total_has_header = True
            except IOError:
                wb_total = Workbook()
                ws_total = wb_total.active
                total_has_header = False

            # 读取 KPI（xls）
            book = xlrd.open_workbook(local_kpi)
            sheet = book.sheet_by_index(0)

            # 表头
            if not total_has_header:
                header = sheet.row_values(0)
                header.extend(["timestamp", "result"])
                ws_total.append(header)

            # 数据行
            for row_idx in range(1, sheet.nrows):
                row = sheet.row_values(row_idx)
                row.extend([ts, result])
                ws_total.append(row)

                fill = pass_fill if result == "pass" else fail_fill
                current_row = ws_total.max_row
                for col in range(1, ws_total.max_column + 1):
                    ws_total.cell(row=current_row, column=col).fill = fill

            wb_total.save(local_total)
            win_sftp.put(local_total, total_kpi_xlsx)

        print("[INFO] KPI 汇总完成 ✔")

        # ======================
        # 关闭连接
        # ======================
        sftp.close()
        ssh.close()
        win_sftp.close()
        win_ssh.close()

        if result == "fail":
            sys.exit(1)

    except Exception as e:
        print("[ERROR]", e)
        sys.exit(1)


if __name__ == "__main__":
    main()
