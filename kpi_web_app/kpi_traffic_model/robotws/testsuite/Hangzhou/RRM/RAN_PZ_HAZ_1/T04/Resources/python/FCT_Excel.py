# -*- coding: utf-8 -*-
"""
:copyright: Nokia Networks
:author: Xu Yuyan
:maintainer: Xu Yuyan
:contact: yuyan.xu@nokia-sbell.com
"""
from datetime import datetime
import openpyxl
import csv
import logging


def get_FCT_info(xls_file, sheet_name, param_list):
    param_value_dic = dict()
    adapt_list = []
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    for param in param_list:
        print('param:', param)
        for i in range(1, sheetname.max_row):
            for j in range(1, sheetname.max_column):
                if (sheetname.cell(i, 1).value == param and param == 'Support Release(s) \\ NIDD adaptationId(s)'):
                    for m in range(2, sheetname.max_column):
                        for n in range(i, sheetname.max_row):
                            if (sheetname.cell(n+1, m).value is not None):
                                tmp_value = sheetname.cell(n+1, m).value.strip()
                                adapt_value = sheetname.cell(i, m).value + ':' + tmp_value
                                adapt_list.append(adapt_value)
                    adapt_list.sort()
                    param_value_dic[param] = adapt_list
                if (sheetname.cell(i, 1).value == param and param != 'Support Release(s) \\ NIDD adaptationId(s)'):
                    param_value_dic[param] = sheetname.cell(i, j+1).value
                    param_match = True
                break
            if (param_match):
                break
        param_match = False
    xls_book.close()
    return param_value_dic


def get_FCT_version(xls_file, sheet_name):
    FTC_version = '0.0.0'
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    start_index = 1
    for i in range(start_index, sheetname.max_row):
        for j in range(1, sheetname.max_column):
            if (sheetname.cell(i, j).value == 'Version'):
                FTC_version = sheetname.cell(i, j+1).value
                param_match = True
                break
        if (param_match):
            break
    start_index = start_index + 1
    xls_book.close()
    return FTC_version


def make_mismatch_FCT(xls_file, sheet_name, param_name, new_value):
    filename = xls_file.split(".")
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    for i in range(1, sheetname.max_row):
        for j in range(1, sheetname.max_column):
            if (sheetname.cell(i, j).value == param_name):
                sheetname.cell(row=i, column=j+1, value=new_value)
                param_match = True
                break
        if (param_match):
            break
    maxRow = sheetname.max_row
    sheetname.delete_rows(maxRow)
    sheetname.delete_rows(maxRow-1)
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def update_parameters(xls_file, sheet_name, param_value_dic):
    filename = xls_file.split(".")
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    for param in param_value_dic.keys():
        param_match = False
        new_value = param_value_dic[param]
        for i in range(1, sheetname.max_row):
            for j in range(1, sheetname.max_column):
                if (sheetname.cell(i, j).value == param):
                    sheetname.cell(row=i, column=j+1, value=new_value)
                    param_match = True
                    break
            if (param_match):
                break
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def add_obsolete(xls_file, sheet_name, new_value_tool, new_value_version):
    filename = xls_file.split(".")
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    obs_list = []
    for i in range(1, sheetname.max_row):
        for j in range(1, sheetname.max_column):
            if (sheetname.cell(i, 1).value != 'Obsoletes tool(s) :'):
                break
            else:
                param_match = True
            if (sheetname.cell(i, 1).value == 'Obsoletes tool(s) :' and sheetname.cell(i, j).value is None):
                sheetname.cell(row=i, column=j, value=new_value_tool)
                sheetname.cell(row=i+1, column=j, value=new_value_version)
                break
        if (param_match):
            break
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def update_with_obsolete(xls_file, sheet_name, param_value_dic, obsolete_tool, obsolete_version):
    filename = xls_file.split(".")
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    obs_list = []
    for param in param_value_dic.keys():
        param_match = False
        new_value = param_value_dic[param]
        for i in range(1, sheetname.max_row):
            for j in range(1, sheetname.max_column):
                if (sheetname.cell(i, j).value == param):
                    sheetname.cell(row=i, column=j+1, value=new_value)
                    param_match = True
                    break
            if (param_match):
                break
    param_match = False
    for i in range(1, sheetname.max_row):
        for j in range(1, sheetname.max_column):
            if (sheetname.cell(i, 1).value != 'Obsoletes tool(s) :'):
                break
            else:
                param_match = True
            if (sheetname.cell(i, 1).value == 'Obsoletes tool(s) :' and sheetname.cell(i, j).value is None):
                sheetname.cell(row=i, column=j, value=obsolete_tool)
                sheetname.cell(row=i+1, column=j, value=obsolete_version)
                break
        if (param_match):
            break
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def update_parameter(xls_file, sheet_name, param_name, new_value):
    filename = xls_file.split(".")
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    for i in range(1, sheetname.max_row):
        for j in range(1, sheetname.max_column):
            if (sheetname.cell(i, j).value == param_name):
                sheetname.cell(row=i, column=j+1, value=new_value)
                param_match = True
                break
        if (param_match):
            break
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def delete_two_lines_adapt(xls_file, sheet_name):
    filename = xls_file.split(".")
    print(filename[0])
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    maxRow = sheetname.max_row
    sheetname.delete_rows(maxRow)
    sheetname.delete_rows(maxRow-1)
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def add_two_lines_adapt(xls_file, sheet_name, row_data_dic):
    filename = xls_file.split(".")
    print(filename[0])
    now = datetime.now()
    time_string = now.strftime("%Y-%m-%d_%H-%M-%S")
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    maxRow = sheetname.max_row
    for key in row_data_dic.keys():
        rowdata_list = [key, row_data_dic[key]]
        sheetname.append(rowdata_list)
    xls_book.save(filename[0]+'_updated_' + time_string + '.xlsx')
    return filename[0]+'_updated_' + time_string + '.xlsx'


def get_default_value_ADDIN(xls_file, sheet_name):
    addin_list = []
    xls_book = openpyxl.load_workbook(xls_file)
    sheetname = xls_book[sheet_name]
    param_match = False
    p_list = ['MOC Name', 'Abbreviation', 'Parameter Name', 'Default Value', 'Parent Parameter']
    for i in range(2, sheetname.max_row):
        param_value_dic = dict()
        if (sheetname.cell(i, 1).value == 'ADDIN'):
            for param in p_list:
                for j in range(2, sheetname.max_column):
                    if (sheetname.cell(1, j).value == param):
                        param_value_dic[param] = sheetname.cell(i, j).value
            addin_list.append(param_value_dic)
    xls_book.close()
    return addin_list


def get_imported_additional_csv(csv_file,  prefix='pdcchIntent'):
    with open(csv_file, 'r') as csvfile:
        reader = csv.reader(csvfile)
        csv_data_dic = dict()
        rows = [row for row in reader]
        # print(rows)
        lines_length = len(rows)
        parameter_names = rows[0]
        col_num = len(parameter_names)
        for i in range(1, lines_length):
            one_MO_data = rows[i]
            param_value_dic = dict()
            object_DN = one_MO_data[0]
            for j in range(1, col_num):
                param_names = parameter_names[j].split(".")
                if (param_names[0] == prefix):
                    param_value_dic[param_names[1]] = one_MO_data[j]
            csv_data_dic[object_DN] = param_value_dic
        return csv_data_dic


def check_columns_not_empty_csv(csv_file):
    """
    Check audit report saved from NetAct. its columns: DN, Parameter abbreviated name, UI value, UI value (current),
    at least have one value is not empty for every column.
    """
    check_column = [0, 2, 3, 4]
    not_empty = False
    with open(csv_file, 'r') as csvfile:
        reader = csv.reader(csvfile)
        rows = [row for row in reader]
        lines_length = len(rows)
        for i in range(1, lines_length):
            row = rows[i]
            column_len = len(row)
            # print("debug-len:",column_len)
            for j in range(0, column_len):
                if (row[j] != '' and j in check_column):
                    # print("before:",check_column)
                    check_column.remove(j)
                    # print("after:",check_column)
            if (len(check_column) == 0):
                not_empty = True
                break
        if not (not_empty):
            for e in check_column:
                logging.error("All values in column:" + rows[0][e] + " are empty ")
    return not_empty


if __name__ == "__main__":
    get_FCT_info("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "abc")
    get_FCT_version("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA")
    add_obsolete("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "PDCCH", "4")
    update_parameter("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "Version")
    delete_two_lines_adapt("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA")
    update_parameters("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "Version")
    add_two_lines_adapt("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "abc")
    make_mismatch_FCT("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "Version", "4.2.2")
    update_with_obsolete("FCT_PDCCH_MultiRelease.xlsx", "GENERAL-METADATA", "Version", "PDCCH", "4")
    get_default_value_ADDIN("FCT_PDCCH_MultiRelease.xlsx", "PARAMETER-METADATA")
    get_imported_additional_csv("additionalInfo-1816.csv")
    check_columns_not_empty_csv("audit_report.csv")
