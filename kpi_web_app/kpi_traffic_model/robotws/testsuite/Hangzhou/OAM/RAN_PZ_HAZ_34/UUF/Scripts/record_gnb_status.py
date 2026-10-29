"""
Author: Yafang Chen
Contacts: yafang.chen@nokia-sbell.com
Description: excel writer and reader

"""

import openpyxl
import os
import re
from openpyxl.styles.colors import BLUE, BLACK, COLOR_INDEX
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter

g_status_title = ['Test Run', 'Mem total', 'Mem used', 'Mem free', 'Mem shared', 'Mem buff/cache', 'Mem available',
                  'Swap total', 'Swap used', 'Swap free', 'run space', 'stb space']
thin = Side(border_style="thin", color=BLACK)
g_border = Border(top=thin, left=thin, right=thin, bottom=thin)
g_title_fill = PatternFill(fill_type='solid', fgColor=COLOR_INDEX[5])  # YELLOW = COLOR_INDEX[5]


def write_gnb_status_info_to_excel(test_run, type, info, file_name):
    """
        param:
            test_run: the test run name
            type: could be memory, run space, stb space, or any other names.
            info: the value to fill in
            file_name: the excel file path, no need to create first.
        example:
            | [1]NetAct_C2C_UUF_Upgrade | stb space | 1.9G | /tmp/gnb_status_check.xlsx |
    """
    check_excel_file_exist(file_name)
    wb = openpyxl.load_workbook(file_name)
    ws = wb.active
    r_index = __filter_row_by_test_run(ws, test_run)
    info = info.strip()
    if type == 'memory':
        values = format_memory_info(info)
        for i, v in enumerate(values):
            ws.cell(r_index, i+2).value = v
            ws.cell(r_index, i+2).border = g_border
    elif type == 'run space':
        ws.cell(r_index, 11).value = info
        ws.cell(r_index, 11).border = g_border
    elif type == 'stb space':
        ws.cell(r_index, 12).value = info
        ws.cell(r_index, 12).border = g_border
    else:
        c_index = __filter_column_by_info_type(ws, type)
        ws.cell(r_index, c_index).value = info
        ws.cell(r_index, c_index).border = g_border
        if ws.column_dimensions[get_column_letter(c_index)].width < len(info):
            ws.column_dimensions[get_column_letter(c_index)].width = len(info)
    wb.save(file_name)


def format_memory_info(info):
    info_list = info.strip().split('\r\n')
    men_grep = re.match(r'Mem:\s*(\S*)\s*(\S*)\s*(\S*)\s*(\S*)\s*(\S*)\s*(\S*)\s*', info_list[1])
    swap_grep = re.match(r'Swap:\s*(\S*)\s*(\S*)\s*(\S*)\s*', info_list[2])
    values = [men_grep.group(1), men_grep.group(3), men_grep.group(3), men_grep.group(4), men_grep.group(5),
              men_grep.group(6), swap_grep.group(1), swap_grep.group(2), swap_grep.group(3)]
    return values


def check_excel_file_exist(file_name, only_check=False):
    if not os.path.exists(file_name):
        if only_check:
            return False
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            for i, v in enumerate(g_status_title):
                ws.cell(1, i+1).value = v
                ws.cell(1, i+1).fill = g_title_fill
                ws.cell(1, i+1).border = g_border
                ws.column_dimensions[get_column_letter(i+1)].width = len(v)
            ws.column_dimensions['A'].width = 28.0
            wb.save(file_name)
    else:
        return True


def __filter_row_by_test_run(ws_handler, test_run):
    max_row = ws_handler.max_row
    for i in range(2, max_row + 1):
        if ws_handler.cell(row=i, column=1).value == test_run:
            r_index = i
            break
    else:
        r_index = max_row + 1
        ws_handler.cell(row=r_index, column=1).value = test_run
        ws_handler.cell(row=r_index, column=1).border = g_border
    return r_index


def __filter_column_by_info_type(ws_handler, info_type):
    max_col = ws_handler.max_column
    for i in range(1, max_col + 1):
        if ws_handler.cell(row=1, column=i).value == info_type:
            c_index = i
            break
    else:
        c_index = max_col + 1
        ws_handler.cell(row=1, column=c_index).value = info_type
        ws_handler.cell(row=1, column=c_index).border = g_border
        ws_handler.cell(row=1, column=c_index).fill = g_title_fill
        ws_handler.column_dimensions[get_column_letter(c_index)].width = len(info_type)
    return c_index


if __name__ == '__main__':
    memory = '              total        used        free      shared  buff/cache   available\r\n' \
             'Mem:            16G        8.3G        3.8G        2.1G        3.2G        4.5G\r\n' \
             'Swap:            0B          0B          0B\r\n'
    filename = r'C:\cyf\tmp\gnb_status_check.xlsx'
    write_gnb_status_info_to_excel('[1]NetAct_C2C_UUF_Upgrade', 'memory', memory, filename)
    write_gnb_status_info_to_excel('[1]NetAct_C2C_UUF_Upgrade', 'stb space', '1.9G', filename)
    write_gnb_status_info_to_excel('[1]NetAct_C2C_UUF_Upgrade', 'run space', '1.8G', filename)
    write_gnb_status_info_to_excel('[2]NetAct_C2C_UUF_Upgrade', 'other info', 'test2324', filename)
    write_gnb_status_info_to_excel('[2]NetAct_C2C_UUF_Upgrade', 'memory', memory, filename)
    write_gnb_status_info_to_excel('[3]NetAct_C2C_UUF_Upgrade', 'stb space', '2.0G', filename)
