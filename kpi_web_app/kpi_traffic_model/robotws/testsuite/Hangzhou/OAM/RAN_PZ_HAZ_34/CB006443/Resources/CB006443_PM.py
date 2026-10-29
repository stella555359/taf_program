# -*- coding: utf-8 -*
from taf.transport.ssh import ssh
import logging
import time
from robot.libraries.OperatingSystem import OperatingSystem
from robot.libraries.BuiltIn import BuiltIn
import os
import openpyxl
import xlrd

ssh = ssh()


def get_file_names(directory, connection, file_indicator='*'):

    cmd_list_file_name = "cd {};ls {}".format(directory, file_indicator)
    output = ssh.execute(cmd_list_file_name, connection=connection).strip()
    logging.debug("cmd:{}, output:{}".format(cmd_list_file_name, output))
    list_file_names = []
    if output.find('No such file or directory') == -1:
        list_file_names = output.split()
    else:
        logging.info("No file found under folder {}".format(directory))
    return list_file_names


def is_file_empty(file_path, connection):
    is_file_exists = ssh.file_exists(
        remote_path=file_path,
        connection=connection)
    if is_file_exists:
        cmd_check_if_file_not_empty = "test -s {} && echo True||echo False".format(
            file_path)
        is_file_not_empty = ssh.execute(
            command=cmd_check_if_file_not_empty,
            connection=connection).strip()
        if is_file_not_empty == "True":
            return False
        else:
            return True
    else:
        logging.error(
            msg="File {} does not exist on {}".format(
                file_path, connection))
        return True


def read_pm_parameter_from_nidd(
        filelocation, sheet, MOCName, filter, exception_list=[]):
    """This KW is just for 5gClassicalBTS

    """
    os_obj = OperatingSystem()
    parameter_list = []
    workbook = xlrd.open_workbook(filelocation)
    worksheet = workbook.sheet_by_name(sheet)
    print('excel file is: {}'.format(worksheet.nrows))
    for row in range(6, worksheet.nrows):
        if worksheet.cell_value(row, 1) == MOCName and worksheet.cell_value(row, 29) != 'Not modifiable' \
                and filter not in worksheet.cell_value(row, 18) \
                and worksheet.cell_value(row, 3) not in exception_list:
            parameter_list.append(worksheet.cell_value(row, 3))
    print("parameter list is: {}".format(parameter_list))
    return parameter_list


"""
def read_pm_parameter_from_nidd(filelocation,sheet, MOCName):
####this function is implemented with openyxl,but finish this execution need more time than with package xlrd
    os_obj = OperatingSystem()
    parameter_list = []
    workbook = openpyxl.load_workbook(filelocation)
    worksheet = workbook[sheet]
    print('excel file is: {}'.format(worksheet.max_row))
    for row in range(6, worksheet.max_row + 1):
        if list(worksheet.rows)[row][1].value == MOCName:
            parameter_list.append(list(worksheet.rows)[row][3].value);
            print("parameter list is: {}".format(parameter_list));
    print("parameter list is: {}".format(parameter_list));
    return parameter_list
"""

"""
#####function read_pm_parameter_from_nidd is based on pandas, but this package is not permitted
to install in UTE due to very big size
def read_pm_parameter_from_nidd(filelocation,sheet, MOCName, ignorerows=5):
    os_obj = OperatingSystem()
    parameter_list = []
    workbook = pandas.DataFrame(pandas.read_excel(filelocation, usecols=['MOC Name','Abbreviation'],
                                                  sheet_name=sheet, skiprows=ignorerows))
    print("excel file is: {}".format(workbook))
    worksheet = workbook[workbook['MOC Name'].isin([MOCName])]
    print("excel file is: {}".format(worksheet))
    print("len is: {}".format(worksheet.shape[0]))
    for i in range(worksheet.shape[0]):
        parameter_list.append(worksheet.iloc[i,1]);
    print("parameter list is: {}".format(parameter_list))
    return parameter_list
"""
"""
if __name__ == "__main__":
    # MOCName='NRBTS/NRPMRNL'  MOCName='MNL/MNLENT/PMMNL'
    # MOCName='TNLSVC/PMTNLINT'
    read_pm_parameter_from_nidd(filelocation='/home/ute/NIDD_Parameter_Report_PMMNL.xlsm', sheet='Parameters',
                                MOCName='MNL/MNLENT/PMMNL')
"""
