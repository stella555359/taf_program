# pylint: skip-file
# -*- coding: utf-8 -*-
"""
:created on: 05-16-2024

:copyright: Nokia
"""
from datetime import datetime, timedelta
import importlib
import subprocess
import pkgutil
import os
import json
import logging
import warnings
import zipfile
from zipfile import ZipFile
import tarfile
import lzma
import shutil
import sys

# from robot.api import logger as robot_logger

# module_dir = '/home/ute/LiuNing/robotws/testsuite/Hangzhou/resources/RAN_PZ_HAZ_34/python/'
# sys.path.append(module_dir)
# import ERR_WRN_Check as ERR_WRN_Check

import testsuite.Hangzhou.resources.RAN_PZ_HAZ_34.python.ERR_WRN_Check as ERR_WRN_Check

logging.basicConfig(filename='./ERR_WRN_LOG.log',
                    format='[%(asctime)s-%(filename)s-%(levelname)s:%(message)s]',
                    level=logging.DEBUG, filemode='a',
                    datefmt='%Y-%m-%d%I:%M:%S %p')
warnings.filterwarnings('ignore', category=UserWarning)

compressed_list = ['.gz', '.tar', '.zip', '.xz', '.tgz']
http_proxy = '--proxy=http://10.144.1.10:8080'

NR_message_dic = {56198: 'CPRT_CM_CELLMAPPINGREQ_ID',
                  56199: 'CPRT_CM_CELLMAPPINGRESP_ID',
                  56208: 'CPRT_CM_CELLCONFIGUPDATEREQ_ID',
                  56203: 'CPRT_CM_ANTENNACARRIERACTIVATIONREQ_ID',
                  56204: 'CPRT_CM_ANTENNACARRIERACTIVATIONRESP_ID',
                  56349: 'CPCM_CELLCONFIGUPDATEREQ_ID',
                  56350: 'CPCM_CELLCONFIGUPDATERESP_ID'}
LTE_message_dic = {9885: 'PS_CELL_SETUP_REQ_MSG',
                   30317: 'L2_CELL_SETUP_REQ_MSG',
                   9078: 'MAC_RACH_SETUP_RESP_MSG',
                   8995: 'CELLC_CELL_STATE_CHANGE_REQ_MSG',
                   8996: 'CELLC_CELL_STATE_CHANGE_RESP_MSG'}

column_names = [
    'msg_id_hex',
    'msg_id_dec',
    'msg_name',
    'number']
column_names_cell = [
    'cell_id',
    'cell_type',
    'msg_id_hex',
    'msg_id_dec',
    'msg_name',
    'number']

cell_info = {14843905: 'NR',
             14843904: 'NR',
             14843906: 'NR',
             231938: 'LTE'}


def install_package(package):
    try:
        logging.info(f'start install {package}')
        subprocess.call(['pip', 'install', '-U', 'numpy', package, http_proxy])
    except Exception as e:
        raise (AssertionError("Error installing package " + str(e)))


def install_and_import_package(package):
    # print("pandas library installation")
    if not bool(pkgutil.find_loader(package)):
        install_package(package)
    return importlib.import_module(package)


def analyze_syscom_trace(log_path, analyze_path="/home/ute/syscom_trace", keyword='AaTraceRecord',
                         excluded_keywords=None, suffix='.atr'):
    logging.info(f"log path is {log_path} ")
    new_dir = f"{datetime.now()}"
    new_dir = new_dir.replace(" ", "_")
    if os.path.exists(log_path):
        trace_path = os.path.join(analyze_path, new_dir)
        os.makedirs(trace_path)
    else:
        logging.warning(
            f"{log_path} is not found, ERR_WRN will not be analyzed ")
        return
    decompression_path, files_list = get_syscom_trace_from_snapshot(log_path, trace_path, keyword, suffix,
                                                                    excluded_keywords)
    logging.info(f"decompression_path is {decompression_path} ")
    if not os.path.exists(trace_path):
        trace_path = os.path.join(analyze_path, new_dir)
        os.makedirs(trace_path)
    trace_file_path = os.path.join(trace_path, 'Trace_file')
    os.makedirs(trace_file_path)
    for file in files_list:
        folder_path, gz_filename = os.path.split(file)
        trace_path_name = f"{trace_file_path}/{gz_filename}"
        shutil.copy(file, trace_path_name)
    try:
        shutil.rmtree(decompression_path)
    except Exception as e:
        logging.warning(
            f"there is one exception when clear path {decompression_path}: {e}")
    return trace_file_path


def compare_trace_size(new_trace, refer_trace, max_size,
                       gap_size):
    file_size = 0
    max_size = int(max_size)
    gap_size = int(gap_size)
    logging.info(f"refer_trace file is {refer_trace}")
    buffer_size = 2097152
    if os.path.exists(new_trace):
        file_size = os.path.getsize(new_trace)
    else:
        logging.error(f"there is no new syscom trace, please check")
        return 'fail'
    if refer_trace is None:
        logging.info(f"there is no refer syscom, will not compare trace size")
        if file_size <= max_size:
            logging.info(f"new trace file size is {file_size / 1024 / 1024}MB")
        else:
            logging.error(
                f"new trace file size is {file_size / 1024 / 1024}MB")
            return 'fail'
    else:
        refer_file_size = os.path.getsize(refer_trace)
        logging.info(f"refer_trace file size is {refer_file_size}")
        if (max_size - buffer_size) <= file_size <= max_size:
            logging.info(
                f"new trace file size is {file_size / 1024 / 1024}MB and close to {max_size} "
                f"and and refer trace is {refer_file_size / 1024 / 1024}")
        elif file_size <= (refer_file_size - buffer_size) or (refer_file_size + gap_size) < file_size:
            logging.info(f"new trace file size is {file_size / 1024 / 1024}MB "
                         f"and refer trace is {refer_file_size / 1024 / 1024}")
            logging.error(
                f"new trace file size should not less than refer_trace-buffer_size "
                f"or more than refer_file_size + gap_size")
            return 'fail'
        else:
            logging.info(f"new trace file size is {file_size / 1024 / 1024}MB "
                         f"and refer trace is {refer_file_size / 1024 / 1024}")
    if file_size != 0:
        file_size = (file_size / 1024 / 1024)
    return file_size


def get_syscom_trace_from_snapshot(
        log_path, analyze_path, keyword, suffix, excluded_keywords):
    loop_flag = 0
    if not os.path.exists(analyze_path):
        # shutil.rmtree(analyze_path)
        os.makedirs(analyze_path)
    if os.path.exists(log_path):
        if os.path.isdir(log_path):
            for filename in os.listdir(log_path):
                log_file_name = os.path.join(log_path, filename)
                if os.path.isfile(log_file_name):
                    shutil.copy(log_file_name, analyze_path)
        elif os.path.isfile(log_path):
            shutil.copy(log_path, analyze_path)
    else:
        logging.error(f"{log_path} is not found, please check。")
    log_root_path = None
    log_root_path, decompression_path = ERR_WRN_Check.extract_compressed_file(
        analyze_path, log_root_path, loop_flag)
    files_list = []
    for root, dirs, files in os.walk(log_root_path):
        for file in files:
            if excluded_keywords is None:
                if file.endswith(suffix) and keyword in file:
                    files_list.append(os.path.join(root, file))
            else:
                if file.endswith(
                        suffix) and keyword in file and excluded_keywords not in file:
                    files_list.append(os.path.join(root, file))
    logging.info(
        f"expected file is {files_list}.")
    return decompression_path, files_list


def analyze_pm_trace_and_new_trace(log_path, max_size=15728640, gap_size=1048576,
                                   refer_trace=None, analyze_path="/home/ute/syscom_trace", keyword='AaTraceRecord',
                                   excluded_keywords=None, suffix='.atr',
                                   filesSystem='/dev/vda2',
                                   max_used_space_ratio=0.9):
    trace_file_path = analyze_syscom_trace(
        log_path, analyze_path, keyword, excluded_keywords, suffix)
    new_trace_file = None
    refer_trace_file = None
    file_size = 0
    trace_file_pm_1 = None
    new_trace_size = 0
    pm_trace_size = 0
    logging.info(
        f"refer trace path is {refer_trace} and new trace file is {trace_file_path}")
    if trace_file_path is not None and os.path.exists(trace_file_path):
        for root, dirs, files in os.walk(trace_file_path):
            for file in files:
                if file.endswith(suffix) and "_pm_1_" in file:
                    trace_file_pm_1 = os.path.join(root, file)
                if file.endswith(
                        suffix) and keyword in file and "_pm_" not in file:
                    new_trace_file = os.path.join(root, file)

    else:
        logging.error(
            f"trace_file_path {trace_file_path} is not found ,please check")
    if refer_trace is not None and os.path.exists(refer_trace):
        for root, dirs, files in os.walk(refer_trace):
            for file in files:
                if file.endswith(suffix) and "_pm_" not in file:
                    logging.info(f"refer trace is {file}")
                    refer_trace_file = os.path.join(root, file)
    else:
        logging.warning(f"there is no refer trace and please check")
    disk_status = ERR_WRN_Check.check_disk_state(
        filesSystem, max_used_space_ratio)
    if disk_status:
        logging.error(f"no enough disk space, test will be stopped")
    if trace_file_pm_1 is not None and new_trace_file is not None:
        new_trace_size = compare_trace_size(
            new_trace_file, None, max_size, gap_size)
        pm_trace_size = compare_trace_size(
            trace_file_pm_1, refer_trace_file, max_size, gap_size)
        logging.info(f"pm_trace_size is {pm_trace_size}")
    if new_trace_size == 'fail':
        failed_trace_path = add_tag_for_failed_log(trace_file_path)
        trace_file_path = failed_trace_path
    if pm_trace_size == 'fail':
        failed_trace_path = add_tag_for_failed_log(trace_file_path)
        trace_file_path = failed_trace_path
        if refer_trace is not None:
            failed_refer_trace_path = add_tag_for_failed_log(refer_trace)
    clear_old_trace(analyze_path)
    return trace_file_path, new_trace_size, pm_trace_size, disk_status


def analyze_new_trace_loop(log_path, max_size, gap_size,
                           refer_trace=None, analyze_path="/home/ute/syscom_trace", keyword='AaTraceRecord',
                           excluded_keywords="_pm_", suffix='.atr',
                           filesSystem='/dev/vda2',
                           max_used_space_ratio=0.9):
    trace_file_path = analyze_syscom_trace(
        log_path, analyze_path, keyword, excluded_keywords, suffix)
    new_trace_file = None
    refer_trace_file = None
    file_size = 0
    logging.info(
        f"refer trace path is {refer_trace} and new trace file is {trace_file_path}")
    if trace_file_path is not None and os.path.exists(trace_file_path):
        for root, dirs, files in os.walk(trace_file_path):
            for file in files:
                if file.endswith(suffix) and "_pm_" not in file:
                    logging.info(f"new trace file is {file}")
                    new_trace_file = os.path.join(root, file)
    else:
        logging.error(
            f"trace_file_path {trace_file_path} is not found ,please check")
    if refer_trace is not None and os.path.exists(refer_trace):
        for root, dirs, files in os.walk(refer_trace):
            for file in files:
                if file.endswith(suffix) and "_pm_" not in file:
                    logging.info(f"refer trace is {file}")
                    refer_trace_file = os.path.join(root, file)
    else:
        logging.warning(f"there is no refer trace and please check")
    disk_status = ERR_WRN_Check.check_disk_state(
        filesSystem, max_used_space_ratio)
    if disk_status:
        logging.error(f"no enough disk space, test will be stopped")
    if new_trace_file is not None:
        file_size = compare_trace_size(
            new_trace_file, refer_trace_file, max_size, gap_size)
    if file_size == 'fail':
        failed_trace_path = add_tag_for_failed_log(trace_file_path)
        trace_file_path = failed_trace_path
        logging.info(f"failed_trace_path is {trace_file_path}")
        if refer_trace is not None:
            failed_refer_trace_path = add_tag_for_failed_log(refer_trace)
            logging.info(
                f"failed_refer_trace_path is {failed_refer_trace_path}")
    clear_old_trace(analyze_path)
    return trace_file_path, file_size, disk_status


def add_tag_for_failed_log(trace_path):
    folder_path, trace_filename = os.path.split(trace_path)
    new_path_name = f'{folder_path}_Failed'
    os.rename(folder_path, new_path_name)
    failed_trace_path = f'{new_path_name}/{trace_filename}'
    return failed_trace_path


def clear_old_trace(trace_path):
    dir_list = []
    for entry in os.scandir(trace_path):

        if entry.is_dir() and not entry.is_symlink() and 'fail' not in entry.path.lower():
            dir_list.append(entry.path)
    dir_tuple = [(f, os.path.getctime(f)) for f in dir_list]
    dir_tuple_sorted = sorted(dir_tuple, key=lambda x: x[1], reverse=True)
    logging.info(f"clear trace path is {dir_tuple_sorted}")
    if len(dir_tuple_sorted) > 5:
        for i, item in enumerate(dir_tuple_sorted[5:], 5):
            shutil.rmtree(item[0])


def get_expected_messages_Json_file(file_path, start_time,
                                    end_time):
    install_and_import_package('pandas')
    import pandas as pd
    # now = datetime.now(tz=None)
    messages_filter = []
    start_time = datetime.strptime(start_time, '%Y-%m-%d %H:%M:%S.%f')
    start_time = start_time.strftime('%Y-%m-%d %H:%M:%S.%f')
    end_time = datetime.strptime(end_time, '%Y-%m-%d %H:%M:%S.%f')
    end_time = end_time.strftime('%Y-%m-%d %H:%M:%S.%f')
    logging.info(f"start time is {start_time}, end time is {end_time}")
    df_json = pd.read_json(file_path)
    all_messages = df_json.loc[:, 'messages'][0]
    eight_hours = timedelta(hours=8)
    print(f" eight_hours {eight_hours}")
    for message in all_messages:
        m_time = message['Header']['Time']
        m_name = message['Header']['Name']
        m_id = message['Header']['Number']
        m_content = message['Content']
        m_time = m_time.replace("T", " ")
        m_time = datetime.strptime(m_time, '%Y-%m-%d %H:%M:%S.%f')
        m_time = m_time + eight_hours
        m_time = m_time.strftime('%Y-%m-%d %H:%M:%S.%f')
        if end_time >= m_time >= start_time:
            messages_filter.append(message)
    return messages_filter


def get_monitor_messages_from_config(config_path):
    install_and_import_package('pandas')
    monitor_message = []
    if os.path.exists(config_path):
        with open(config_path, 'r') as file:
            contents = file.readlines()
            for line in contents:
                message_id = line.split(' ', 1)[0]
                message_name = line.split(' ', 1)[1]
                message_id = message_id.replace("0x", "")
                message_id_decimal = int(message_id, 16)
                message_name = message_name.replace("#", "")
                message_name = message_name.replace("\n", "")
                message = {message_id_decimal: message_name}
                monitor_message.append(message)
    else:
        logging.warning(f"there is no CNI-103148 trace configration")
    return monitor_message


def analyze_json_trace_with_cell(
        message_filter, message_analyse_excel, cell_info_dic=None):
    if cell_info_dic is None:
        cell_info_dic = cell_info
    install_and_import_package('pandas')
    import pandas as pd
    install_and_import_package('openpyxl')
    import openpyxl
    message_miss = []
    message_found = []
    message_cell_analyse = pd.DataFrame(columns=column_names_cell)
    for key in cell_info_dic:
        if cell_info_dic[key] == 'NR':
            for monitor_message_id in NR_message_dic:
                num = 0
                msg_id_hex = hex(monitor_message_id)
                for message in message_filter:
                    m_content = message['Content']
                    m_id = message['Header']['Number']
                    m_content_string = json.dumps(m_content)
                    if str(key) in m_content_string and m_id == monitor_message_id:
                        num = num + 1
                if num == 0:
                    logging.error(
                        f" cell {key} miss message {monitor_message_id}: {NR_message_dic[monitor_message_id]}")
                data = [
                    key,
                    'NR',
                    msg_id_hex,
                    monitor_message_id,
                    NR_message_dic[monitor_message_id],
                    num]
                message_cell_analyse.loc[len(
                    message_cell_analyse.index)] = data
        elif cell_info_dic[key] == 'LTE':
            for monitor_message_id in LTE_message_dic:
                num = 0
                msg_id_hex = hex(monitor_message_id)
                for message in message_filter:
                    m_content = message['Content']
                    m_id = message['Header']['Number']
                    m_content_string = json.dumps(m_content)
                    if str(key) in m_content_string and m_id == monitor_message_id:
                        num = num + 1
                if num == 0:
                    logging.error(
                        f" cell {key} miss message {monitor_message_id}: {LTE_message_dic[monitor_message_id]}")
                else:
                    logging.info(
                        f" found message {monitor_message_id}: {LTE_message_dic[monitor_message_id]}, the num is {num}")
                data = [
                    key,
                    'LTE',
                    msg_id_hex,
                    monitor_message_id,
                    LTE_message_dic[monitor_message_id],
                    num]
                message_cell_analyse.loc[len(
                    message_cell_analyse.index)] = data
    workbook = openpyxl.load_workbook(message_analyse_excel)
    if 'cell_message_analyse' in workbook.sheetnames:
        workbook.remove(workbook['cell_message_analyse'])
        workbook.save(message_analyse_excel)
        workbook.close()
    with pd.ExcelWriter(message_analyse_excel, engine='openpyxl', mode='a') as writer:
        message_cell_analyse.to_excel(
            writer, sheet_name='cell_message_analyse', index=False)
    return message_analyse_excel


def analyze_one_json_trace(
        analyze_path, message_filter, monitor_message, message_analyse_excel):
    install_and_import_package('pandas')
    import pandas as pd
    install_and_import_package('openpyxl')
    import openpyxl
    message_miss = []
    message_found = []
    root_path, filename = os.path.split(analyze_path)
    sheet_name = os.path.splitext(filename)[0]
    message_analyse = pd.DataFrame(columns=column_names)
    if len(message_filter) != 0:
        for item in monitor_message:
            monitor_message_id = list(item.keys())[0]
            # print(f" monitor_message_id {monitor_message_id}")
            num = 0
            msg_id_hex = hex(monitor_message_id)
            for message in message_filter:
                m_time = message['Header']['Time']
                m_name = message['Header']['Name']
                m_id = message['Header']['Number']
                m_content = message['Content']
                if m_id == monitor_message_id:
                    num = num + 1
            if num == 0:
                # logging.error(
                #    f" miss message {monitor_message_id}: {item[monitor_message_id]}")
                # print(
                # f" miss message {monitor_message_id}:
                # {item[monitor_message_id]}")
                message_miss.append(item)
                data = [
                    msg_id_hex,
                    monitor_message_id,
                    item[monitor_message_id],
                    num]
                message_analyse.loc[len(message_analyse.index)] = data
            else:
                message = {
                    monitor_message_id: {
                        'name': {
                            item[monitor_message_id]},
                        'num': num}}
                message_found.append(message)
                data = [
                    msg_id_hex,
                    monitor_message_id,
                    item[monitor_message_id],
                    num]
                message_analyse.loc[len(message_analyse.index)] = data
                logging.info(
                    f" found message {monitor_message_id}: {item[monitor_message_id]}, the num is {num}")
                # print(
                # f" found message {monitor_message_id}:
                # {item[monitor_message_id]}, the num is {num}")
            book = openpyxl.load_workbook(message_analyse_excel)
            message_analyse.to_excel(
                message_analyse_excel,
                sheet_name=sheet_name,
                index=False)
    else:
        logging.error(
            f" no monitored messages found ")
    return message_analyse_excel


def analyze_trace_and_generate_result(
        trace_path, config_path, start_time, end_time, suffix, cell_info_dic=None):
    monitor_message = get_monitor_messages_from_config(config_path)
    if len(monitor_message) == 0:
        logging.warning(
            f"stop analyze_trace_and_generate_result due to trace conf missing")
        return None
    current_time = datetime.now()
    timestamp = current_time.strftime('%Y%m%d%H%M%S')
    message_analyse_excel = ERR_WRN_Check.creat_excel_file(
        trace_path, 'message_analyse', timestamp)
    logging.info(
        f" message_analyse_excel is {message_analyse_excel}")
    logging.info(
        f" trace_path is {trace_path}")
    for root, dirs, files in os.walk(trace_path):
        logging.info(
            f" files in trace_path are {files} and suffix is {suffix}")
        for file in files:
            logging.info(
                f" file in trace_path are {file}")
            if file.endswith(suffix):
                json_file = os.path.join(root, file)
                logging.info(
                    f" json file is {json_file}")
                message_filter = get_expected_messages_Json_file(
                    json_file, start_time, end_time)
                message_analyse_excel = \
                    analyze_one_json_trace(
                        json_file,
                        message_filter,
                        monitor_message,
                        message_analyse_excel)
                message_analyse_excel = \
                    analyze_json_trace_with_cell(
                        message_filter, message_analyse_excel, cell_info_dic)


def get_expected_files(input_file, keyword, suffix='.atr'):
    matching_files = []
    matching_files = extract_snapshot_and_get_expected_file(
        input_file, keyword)
    if len(matching_files) == 0:
        logging.info(
            f"there is no expected file in {input_file} please check。")
    return matching_files


def creat_excel_xlsx(log_path, filename, timestamp):
    install_and_import_package('openpyxl')
    import openpyxl
    new_filename = f"{filename}_{timestamp}.xlsx"
    output_dir = os.path.dirname(log_path)
    log_file_name = f"ERR_WRN_LOG_{timestamp}.txt"
    outfile_excel = os.path.join(log_path, new_filename)
    log_file = os.path.join(log_path, log_file_name)
    # print(f"log file is {log_file} ")
    try:
        os.remove(outfile_excel)
    except OSError:
        pass
    wb = openpyxl.Workbook()
    wb.save(outfile_excel)
    return outfile_excel


def extract_snapshot_and_get_expected_file(input_file, keyword, suffix='.atr'):
    file_type = ERR_WRN_Check.get_file_type(input_file)
    folder_path, gz_filename = os.path.split(input_file)
    extracted_filename = os.path.splitext(gz_filename)[0]
    out_file_path = os.path.join(folder_path, extracted_filename)
    expected_files_list = []
    if os.path.exists(input_file) and os.path.getsize(
            input_file) != 0 and file_type == 'zip':
        os.chmod(input_file, 0o777)
        try:
            with zipfile.ZipFile(input_file, 'r') as file_in:
                file_in.extractall(path=out_file_path)
                logging.info(
                    f"file list is {file_in.namelist()}")
        except Exception as e:
            print(f"Error opening zip file: {e}")
        if os.path.exists(out_file_path):
            if os.path.isdir(out_file_path):
                for root, dirs, files in os.walk(out_file_path):
                    # for dir_name in dirs:
                    #    os.chmod(os.path.join(root, dir_name), 0o777)
                    for file_name in files:
                        os.chmod(os.path.join(root, file_name), 0o777)
                        file_type = ERR_WRN_Check.get_file_type(input_file)
                        logging.info(
                            f"file type is {file_type}")
                        file_full_name = os.path.join(root, file_name)
                        loop_flag = 0
                        logging.info(
                            f"file file_full_name is {file_full_name}")
                        expected_files_list = find_expected_files(file_full_name, keyword, loop_flag,
                                                                  expected_files_list, suffix)

        else:
            logging.error(
                f"{out_file_path} is not found")

    else:
        logging.error(
            f"{input_file} is not found, please check。")
    return expected_files_list


def find_expected_files(file_full_name, keyword,
                        loop_flag, expected_files_list, suffix):
    folder_path, gz_filename = os.path.split(file_full_name)
    extracted_filename = os.path.splitext(gz_filename)[0]
    file_suffix = os.path.splitext(gz_filename)[-1]
    logging.info(
        f"file_suffix is {file_suffix}")
    if file_suffix in compressed_list:
        find_in_one_compressed_file(
            file_full_name,
            keyword,
            file_suffix,
            expected_files_list,
            suffix,
            loop_flag)
    else:
        if file_suffix == suffix and keyword in file_full_name:
            expected_files_list.append(file_full_name)
    return expected_files_list


def find_in_one_compressed_file(
        input_file, keyword, file_suffix, expected_files_list, suffix, loop_flag):
    if file_suffix == '.tgz' or file_suffix == '.tar':
        try:
            with tarfile.open(input_file, 'r:*') as file_in:
                for file in file_in.getmembers():
                    logging.info(
                        f"file is {file.name}")
                    file_suffix = check_file_suffix(file.name)
                    file_name = os.path.join(input_file, file.name)
                    if file_suffix in compressed_list:
                        find_expected_files(
                            file_name, keyword, loop_flag, expected_files_list, suffix)
                    else:
                        if file_suffix == suffix and keyword in file.name:
                            expected_files_list.append(file_name)
        except Exception as e:
            logging.info(
                f"An error occurred while extracting {input_file}: {e}")
    if file_suffix == '.gz':
        try:
            with tarfile.open(input_file, 'r:gz') as tar:
                for file in tar.getmembers():
                    logging.info(
                        f".gz file is {file.name}")
                    file_suffix = check_file_suffix(file.name)
                    file_name = os.path.join(input_file, file.name)
                    if file_suffix in compressed_list:
                        find_expected_files(
                            file_name, keyword, loop_flag, expected_files_list, suffix)
                    else:
                        if file_suffix == suffix and keyword in file.name:
                            expected_files_list.append(file_name)
        except Exception as e:
            logging.info(
                f"An error occurred while extracting {input_file}: {e}")
    elif file_suffix == '.zip':
        try:
            with zipfile.ZipFile(input_file, 'r') as file_in:
                for file in file_in.infolist():
                    logging.info(
                        f"file is {file.filename}")
                    file_suffix = check_file_suffix(file)
                    file_name = file_in.read(file.filename)
                    # file_name = os.path.join(input_file, file)
                    if file_suffix in compressed_list:
                        find_expected_files(
                            file_name, keyword, loop_flag, expected_files_list, suffix)
                    else:
                        if file_suffix == suffix and keyword in file:
                            expected_files_list.append(file_name)
        except Exception as e:
            logging.info(f"Error opening zip file: {e}")
    elif file_suffix == '.xz':
        try:
            with lzma.open(input_file, 'rb') as file_in:
                with ZipFile(file_in) as zip_file:
                    for file in zip_file.namelist():
                        logging.info(
                            f"file is {file}")
                        file_suffix = check_file_suffix(file)
                        file_name = os.path.join(input_file, file)
                        if file_suffix in compressed_list:
                            find_expected_files(
                                file_name, keyword, loop_flag, expected_files_list, suffix)
                        else:
                            if file_suffix == suffix and keyword in file:
                                expected_files_list.append(file_name)
        except EOFError:
            logging.warning("There is EOFError")
    else:
        logging.info('uncompressed file end.')


def check_file_suffix(file_full_name):
    folder_path, gz_filename = os.path.split(file_full_name)
    extracted_filename = os.path.splitext(gz_filename)[0]
    file_suffix = os.path.splitext(gz_filename)[-1]
    return file_suffix


if __name__ == "__main__":
    log_path = '/home/ute/syscom_trace/2025-01-06_17:33:29.444367/MRBTS-906_Snapshot_20250106_172952.zip'
    trace_path = '/home/ute/syscom_trace'
    config_path = "/home/ute/syscom_trace/AaTraceRecordFilter.txt"
    clear_old_trace(trace_path)
    # analyze_trace_and_generate_result(
    #    trace_path,
    #    config_path,
    #    suffix='.json',
    #    start_time="2024-12-31 14:17:15.077873",
    #    end_time="2024-12-31 14:21:21.829325")
