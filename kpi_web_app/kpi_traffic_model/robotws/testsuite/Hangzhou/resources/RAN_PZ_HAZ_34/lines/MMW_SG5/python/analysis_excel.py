# -*- coding: utf-8 -*-
# @Time    : 2024/1/9 15:53
# @Author  : Jian.1.Chen
# @Email   : jian.1.chen@nokia-sbell.com
# @File    : analysis_excel.py
# @Software: PyCharm
from openpyxl import load_workbook


def add_new_line_to_target_sheet_for_8132(excel, rsrp, build, time, fiveqi_dict):
    build = build.replace("SBTS", "")
    build = build.replace("_ENB_9999_", "_")
    build = build.replace("0000", "")
    build = build + "_0" + str(time)
    dl_line = [rsrp, build]
    ul_line = [rsrp, build]
    if "RLC" in next(iter(fiveqi_dict)):
        ul_sheet = "Uplink RLC delay"
        dl_sheet = "Downlink RLC delay"
    else:
        ul_sheet = "Uplink PDCP delay"
        dl_sheet = "Downlink PDCP delay"

    for key in fiveqi_dict.keys():
        if "DL" in key:
            dl_line.append(round(float(fiveqi_dict[key]), 3))
        elif "UL" in key:
            ul_line.append(round(float(fiveqi_dict[key]), 3))
    add_new_line_to_target_sheet(excel, ul_sheet, ul_line, rsrp)
    add_new_line_to_target_sheet(excel, dl_sheet, dl_line, rsrp)


def add_new_line_to_target_sheet(excel, target_sheet, new_line, rsrp):
    wb = load_workbook(excel)
    wbs = wb[target_sheet]
    total_rows = wbs.max_row + 1
    for row in range(1, total_rows):
        col_a_cell = "A" + str(row)
        col_b_cell = "B" + str(row)
        if rsrp == wbs[col_a_cell].value and new_line[1] == wbs[col_b_cell].value:
            time = int(new_line[1][-2:-1]) + 1
            new_line[1] = new_line[1][:-2] + str(time) + new_line[1][-1]
    wbs.append(new_line)
    wb.save(excel)


if __name__ == "__main__":
    a = {"5QI-1": "1", "5QI-2": "1"}
    add_new_line_to_target_sheet("C:\\N-20S1PF2R6FZ0-Data\\jian\\Downloads\\test_result_report_template.xlsx",
                                 "Uplink RLC delay", "a", "b")
