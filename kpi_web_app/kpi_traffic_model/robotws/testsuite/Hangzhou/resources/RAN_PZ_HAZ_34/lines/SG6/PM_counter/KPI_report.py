# -*- coding: utf-8 -*-
"""
:copyright: Nokia Networks
:author: Chen Jun
:contact: jun.8.chen@nokia-sbell.com
"""
import logging
from collections import defaultdict
import os
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.chart import (
    LineChart,
    Reference,
    BarChart,
    Series
)
from taf.rfrl.utilities.reports.excel_report import ExcelReport


def get_group_name(report):
    KPIs = report[0]["report"]
    group_name_list = []
    for element in KPIs:
        group_name = element["group_name"]
        if group_name not in group_name_list:
            group_name_list.append(group_name)
        else:
            pass
    return group_name_list


def create_workbook_sheets(group_name_list, file_path, sheet='Sheet'):
    if os.path.exists(file_path):
        os.remove(file_path)
    wb = Workbook()
    wb.remove(wb[sheet])
    for group_name in group_name_list:
        ws = wb.create_sheet(group_name)
    wb.save(file_path)


def get_data(report, group_name_list):
    kpi_report = defaultdict(list)
    for group_name in group_name_list:
        for element in report[0]["report"]:
            name = element["name"]
            value = element["global_result"]["value"]
            rounded_value = round_to_2decimal(value)
            code = element["code"]
            unit = element.get("unit")
            grp_name = element["group_name"]
            kpi = {}
            kpi["code"] = code
            kpi["name"] = name
            kpi["value"] = rounded_value
            kpi["unit"] = unit
            if grp_name == group_name:
                kpi_report[group_name].append(kpi)
            else:
                pass
    return kpi_report


def round_to_2decimal(value):
    try:
        return round(value, 2)
    except TypeError:
        return "N/A"


def write_data_to_xlsx(file_path, kpi_report, start_time, end_time):
    value_name = "Value_{}_{}".format(start_time, end_time)
    report_keys = kpi_report.keys()
    for key in report_keys:
        workbook = load_workbook(file_path)
        sheet = workbook[key]
        sheet["A1"] = "Code"
        sheet["B1"] = "KPI_name"
        sheet["C1"] = "Unit"
        sheet["D1"] = value_name
        row = 2
        for element in kpi_report[key]:
            code = element["code"]
            name = element["name"]
            unit = element["unit"]
            value = element["value"]
            row = int(row)
            sheet["A%d" % (row)] = code
            sheet["B%d" % (row)] = name
            sheet["C%d" % (row)] = unit
            sheet["D%d" % (row)] = value
            row = row + 1
        workbook.save(filename=file_path)
        workbook.close()


def append_col_to_xlsx(file_path, kpi_report, start_time, end_time, location):
    value_name = "Value_{}_{}".format(start_time, end_time)
    x = "{}1".format(location)
    report_keys = kpi_report.keys()
    for key in report_keys:
        workbook = load_workbook(file_path)
        sheet = workbook[key]
        sheet[x] = value_name
        row = 2
        for element in kpi_report[key]:
            code = element["code"]
            print("cur is {}".format(code))
            value = element["value"]
            row = int(row)
            code_name = sheet["A%d" % (row)].value
            print("origin is {}".format(code_name))
            if code == code_name:
                y = "{}%d".format(location)
                sheet[y % (row)] = value
            row = row + 1
        workbook.save(filename=file_path)
        workbook.close()


def get_col_num(file_path, sheet):
    mapping = {1: "A", 2: "B", 3: "C", 4: "D", 5: "E", 6: "F", 7: "G", 8: "H"}
    workbook = load_workbook(file_path)
    sheet = workbook[sheet]
    col_num_max = sheet.max_column
    print("col_num_max is {}".format(col_num_max))
    col_num_max = col_num_max + 1
    col_no = mapping[col_num_max]
    return col_no


def generate_data_report_from_aasyslog(tput_data_list, header_list, file_path):
    """
    author: lei.jia@nokia-sbell.com
    """
    file_dir = os.path.dirname(file_path)
    if not os.path.exists(file_dir):
        os.mkdir(file_dir)
    collective_report = ExcelReport(excel_file_path=file_path)
    if not collective_report.file_exists:
        collective_report.create_new_excel_file(headers=header_list, data=tput_data_list)
    else:
        collective_report.add_data_to_worksheet(data=tput_data_list)
