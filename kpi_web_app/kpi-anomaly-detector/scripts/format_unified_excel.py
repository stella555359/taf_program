#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Format kpi_config_unified.xlsx with proper styling

Features:
- Freeze top row
- Header with colors and wrap text
- Override_Strategy, Enabled columns with dropdowns
- Physical_Constraint column for dB/percentage constraints
- Remove Forbidden_Methods column
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from pathlib import Path


def format_unified_excel():
    """Format the unified Excel file with proper styling"""
    
    lib_dir = Path(__file__).parent.parent / 'lib' / 'kpi_xml'
    input_file = lib_dir / 'kpi_config_unified.xlsx'
    output_file = input_file  # Overwrite
    
    print(f"Loading: {input_file}")
    df = pd.read_excel(input_file)
    
    # Reorder columns: move Override_Strategy and Enabled after Detection_Strategy
    cols = list(df.columns)
    
    # Remove columns to reposition
    if 'Override_Strategy' in cols:
        cols.remove('Override_Strategy')
    if 'Enabled' in cols:
        cols.remove('Enabled')
    if 'Forbidden_Methods' in cols:
        cols.remove('Forbidden_Methods')  # Remove this column
    
    # Find Detection_Strategy position and insert after it
    det_idx = cols.index('Detection_Strategy')
    cols.insert(det_idx + 1, 'Override_Strategy')
    cols.insert(det_idx + 2, 'Enabled')
    
    # Add Physical_Constraint column (new)
    if 'Physical_Constraint' not in cols:
        cols.insert(det_idx + 3, 'Physical_Constraint')
        df['Physical_Constraint'] = ''
    
    # Ensure columns exist
    if 'Override_Strategy' not in df.columns:
        df['Override_Strategy'] = ''
    if 'Enabled' not in df.columns:
        df['Enabled'] = False
    
    df = df[cols]
    
    # Save with formatting
    print(f"Saving: {output_file}")
    df.to_excel(output_file, index=False, sheet_name='1_All_Codes_Classification')
    
    # Apply formatting with openpyxl
    wb = load_workbook(output_file)
    ws = wb.active
    
    # Colors
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    override_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')  # Orange
    enabled_fill = PatternFill(start_color='92D050', end_color='92D050', fill_type='solid')  # Green
    constraint_fill = PatternFill(start_color='00B0F0', end_color='00B0F0', fill_type='solid')  # Blue
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Format header row
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        
        # Special colors for override columns
        if col_name == 'Override_Strategy':
            cell.fill = override_fill
            cell.font = Font(bold=True, color='000000')
        elif col_name == 'Enabled':
            cell.fill = enabled_fill
            cell.font = Font(bold=True, color='000000')
        elif col_name == 'Physical_Constraint':
            cell.fill = constraint_fill
            cell.font = Font(bold=True, color='000000')
        else:
            cell.fill = header_fill
    
    # Set row height for header
    ws.row_dimensions[1].height = 30
    
    # Freeze top row
    ws.freeze_panes = 'A2'
    
    # Find column letters for dropdowns
    override_col = get_column_letter(df.columns.get_loc('Override_Strategy') + 1)
    enabled_col = get_column_letter(df.columns.get_loc('Enabled') + 1)
    constraint_col = get_column_letter(df.columns.get_loc('Physical_Constraint') + 1)
    
    # Strategy dropdown
    strategy_list = 'ABS_THRESHOLD,IQR,POISSON,RELATIVE,STANDARD_SPC,Z_SCORE,(use default)'
    strategy_dv = DataValidation(type='list', formula1=f'"{strategy_list}"', allow_blank=True)
    strategy_dv.error = 'Please select a valid strategy'
    strategy_dv.errorTitle = 'Invalid Strategy'
    ws.add_data_validation(strategy_dv)
    strategy_dv.add(f'{override_col}2:{override_col}{len(df)+1}')
    
    # Enabled dropdown
    enabled_dv = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(enabled_dv)
    enabled_dv.add(f'{enabled_col}2:{enabled_col}{len(df)+1}')
    
    # Physical_Constraint dropdown
    constraint_list = '1dB floor,percentage floor,count floor,ratio check,(none)'
    constraint_dv = DataValidation(type='list', formula1=f'"{constraint_list}"', allow_blank=True)
    ws.add_data_validation(constraint_dv)
    constraint_dv.add(f'{constraint_col}2:{constraint_col}{len(df)+1}')
    
    # Auto-width columns (approximate)
    for col_idx, col_name in enumerate(df.columns, 1):
        max_len = max(len(str(col_name)), 10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 25)
    
    # Special width for certain columns
    ws.column_dimensions['C'].width = 40  # Name
    ws.column_dimensions[override_col].width = 18
    ws.column_dimensions[enabled_col].width = 10
    ws.column_dimensions[constraint_col].width = 18
    
    wb.save(output_file)
    
    print(f"\n✅ Formatted Excel saved!")
    print(f"   Columns: {len(df.columns)}")
    print(f"   Rows: {len(df)}")
    print(f"\nNew column order around Detection_Strategy:")
    det_idx = list(df.columns).index('Detection_Strategy')
    for i in range(det_idx, min(det_idx + 4, len(df.columns))):
        print(f"   [{i+1}] {df.columns[i]}")


if __name__ == '__main__':
    format_unified_excel()
