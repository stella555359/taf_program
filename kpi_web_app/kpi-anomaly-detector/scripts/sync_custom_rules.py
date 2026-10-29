#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Sync Custom Rules from Classification and Comparison Reports

This script:
1. Reads KPI types and strategies from kpi_classification_v2.xlsx
2. Updates kpi_custom_rules.xlsx with classification data
3. Marks "Other→Regression" KPIs for user attention (Enabled=FALSE by default)

User workflow:
1. Run this script to populate rules
2. Open kpi_custom_rules.xlsx
3. Filter by "Change_Type" = "Other→Regression" (these need attention)
4. Review and set Enabled=TRUE for KPIs that should use custom strategy
5. Run: python scripts/manage_custom_rules.py --apply
"""

import sys
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Paths
BASE_DIR = Path(__file__).parent.parent
LIB_DIR = BASE_DIR / 'lib' / 'kpi_xml'
REPORTS_DIR = BASE_DIR / 'reports'

CLASSIFICATION_FILE = LIB_DIR / 'kpi_classification_v2.xlsx'
CUSTOM_RULES_FILE = LIB_DIR / 'kpi_custom_rules.xlsx'
COMPARISON_FILE = REPORTS_DIR / 'kpi_change_comparison.xlsx'
HISTORY_FILE = LIB_DIR / 'kpi_rules_history.json'

STRATEGIES = [
    "STANDARD_SPC",
    "ABS_THRESHOLD",
    "RANGE_CONTEXT",
    "RELATIVE_CHANGE",
    "ABS_DELTA_TREND",
    "QUANTILES",
    "TAIL_AWARE",
    "WINDOW_COUNT",
    "SEMANTIC_RULE",
    "(使用默认)"
]


def load_classification_data():
    """Load KPI classification from kpi_classification_v2.xlsx"""
    print(f"Loading classification from: {CLASSIFICATION_FILE}")
    
    df = pd.read_excel(CLASSIFICATION_FILE, sheet_name='1_All_Codes_Classification')
    
    classification = {}
    for _, row in df.iterrows():
        code = str(row.get('Code', '')).strip()
        if not code:
            continue
        
        classification[code] = {
            'KPI_Name': str(row.get('Name', '')).strip(),
            'KPI_Type': str(row.get('KPI_Type', 'unknown')).strip(),
            'Default_Strategy': str(row.get('Detection_Strategy', 'STANDARD_SPC')).strip(),
            'Units': str(row.get('Units', '')).strip(),
            'Level': str(row.get('Level', '')).strip(),
        }
    
    print(f"  Loaded {len(classification)} KPIs from classification")
    return classification


def load_comparison_data():
    """Load Other→Regression KPIs from comparison report"""
    if not COMPARISON_FILE.exists():
        print(f"Warning: Comparison file not found: {COMPARISON_FILE}")
        return set(), set()
    
    print(f"Loading comparison from: {COMPARISON_FILE}")
    
    # Load Other→Regression (degraded)
    try:
        df_degraded = pd.read_excel(COMPARISON_FILE, sheet_name='3_Other_to_Regression')
        degraded_codes = set(str(row['Code']).strip() for _, row in df_degraded.iterrows() if row.get('Code'))
        print(f"  Found {len(degraded_codes)} Other→Regression KPIs")
    except Exception as e:
        print(f"  Warning: Could not load 3_Other_to_Regression: {e}")
        degraded_codes = set()
    
    # Load Regression→Other (improved)
    try:
        df_improved = pd.read_excel(COMPARISON_FILE, sheet_name='2_Regression_to_Other')
        improved_codes = set(str(row['Code']).strip() for _, row in df_improved.iterrows() if row.get('Code'))
        print(f"  Found {len(improved_codes)} Regression→Other KPIs")
    except Exception as e:
        print(f"  Warning: Could not load 2_Regression_to_Other: {e}")
        improved_codes = set()
    
    return degraded_codes, improved_codes


def load_existing_rules():
    """Load existing user edits from kpi_custom_rules.xlsx"""
    if not CUSTOM_RULES_FILE.exists():
        return {}
    
    try:
        df = pd.read_excel(CUSTOM_RULES_FILE, sheet_name='1_Custom_Rules')
        existing = {}
        for _, row in df.iterrows():
            code = str(row.get('Code', '')).strip()
            if not code:
                continue
            
            override = str(row.get('Override_Strategy', '')).strip()
            enabled = str(row.get('Enabled', 'FALSE')).strip()
            comment = str(row.get('Comment', '')).strip()
            
            # Only preserve if user has made edits
            if enabled.upper() == 'TRUE' or (override and override != '(使用默认)') or comment:
                existing[code] = {
                    'Override_Strategy': override,
                    'Enabled': enabled,
                    'Comment': comment
                }
        
        print(f"  Preserved {len(existing)} existing user rules")
        return existing
    except Exception as e:
        print(f"  Warning: Could not load existing rules: {e}")
        return {}


def create_custom_rules_excel(classification, degraded_codes, improved_codes, existing_rules):
    """Create updated kpi_custom_rules.xlsx"""
    
    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    edit_fill = PatternFill(start_color="FFFFD7", end_color="FFFFD7", fill_type="solid")  # Yellow
    degraded_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Red
    improved_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
    attention_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Orange
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    wb = Workbook()
    
    # Sheet 1: Custom Rules
    ws_rules = wb.active
    ws_rules.title = "1_Custom_Rules"
    
    headers = ['Code', 'KPI_Name', 'KPI_Type', 'Default_Strategy', 'Override_Strategy', 
               'Enabled', 'Change_Type', 'Priority', 'Comment']
    for col, h in enumerate(headers, 1):
        cell = ws_rules.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    # Data validation
    dv_strategy = DataValidation(type="list", formula1=f'"{",".join(STRATEGIES)}"', allow_blank=True)
    dv_strategy.error = "请从下拉列表中选择策略"
    ws_rules.add_data_validation(dv_strategy)
    
    dv_enabled = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws_rules.add_data_validation(dv_enabled)
    
    # Sort: Degraded first, then by code
    all_codes = sorted(classification.keys())
    
    # Reorder: degraded first, then improved, then others
    degraded_list = sorted([c for c in all_codes if c in degraded_codes])
    improved_list = sorted([c for c in all_codes if c in improved_codes and c not in degraded_codes])
    other_list = sorted([c for c in all_codes if c not in degraded_codes and c not in improved_codes])
    
    ordered_codes = degraded_list + improved_list + other_list
    
    row_idx = 2
    for code in ordered_codes:
        info = classification.get(code, {})
        existing = existing_rules.get(code, {})
        
        # Determine change type and priority
        if code in degraded_codes:
            change_type = "Other→Regression ⚠️"
            priority = "HIGH - 需要关注"
        elif code in improved_codes:
            change_type = "Regression→Other ✓"
            priority = "LOW - 已改善"
        else:
            change_type = ""
            priority = ""
        
        # Use existing values if available, otherwise defaults
        override = existing.get('Override_Strategy', '')
        enabled = existing.get('Enabled', 'FALSE')
        comment = existing.get('Comment', '')
        
        # Write data
        ws_rules.cell(row=row_idx, column=1, value=code).border = thin_border
        ws_rules.cell(row=row_idx, column=2, value=info.get('KPI_Name', '')).border = thin_border
        ws_rules.cell(row=row_idx, column=3, value=info.get('KPI_Type', '')).border = thin_border
        ws_rules.cell(row=row_idx, column=4, value=info.get('Default_Strategy', '')).border = thin_border
        
        # Override_Strategy - editable
        cell = ws_rules.cell(row=row_idx, column=5, value=override)
        cell.fill = edit_fill
        cell.border = thin_border
        dv_strategy.add(cell)
        
        # Enabled - editable
        cell = ws_rules.cell(row=row_idx, column=6, value=enabled)
        cell.fill = edit_fill
        cell.border = thin_border
        dv_enabled.add(cell)
        
        # Change_Type - highlight
        cell = ws_rules.cell(row=row_idx, column=7, value=change_type)
        cell.border = thin_border
        if code in degraded_codes:
            cell.fill = degraded_fill
        elif code in improved_codes:
            cell.fill = improved_fill
        
        # Priority
        cell = ws_rules.cell(row=row_idx, column=8, value=priority)
        cell.border = thin_border
        if code in degraded_codes:
            cell.fill = attention_fill
        
        # Comment - editable
        cell = ws_rules.cell(row=row_idx, column=9, value=comment)
        cell.border = thin_border
        
        row_idx += 1
    
    # Column widths
    col_widths = [18, 55, 18, 18, 18, 10, 22, 18, 35]
    for i, w in enumerate(col_widths, 1):
        ws_rules.column_dimensions[chr(64+i)].width = w
    
    ws_rules.freeze_panes = 'A2'
    ws_rules.auto_filter.ref = f"A1:I{row_idx-1}"
    
    # Sheet 2: Strategy Reference
    ws_ref = wb.create_sheet(title="2_Strategy_Reference")
    ws_ref['A1'] = "检测策略说明"
    ws_ref['A1'].font = Font(bold=True, size=14)
    
    ref_headers = ['Strategy', 'Description', 'Suitable_For', 'When_to_Override']
    for col, h in enumerate(ref_headers, 1):
        cell = ws_ref.cell(row=3, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    ref_data = [
        ("STANDARD_SPC", "标准SPC控制图方法", "通用，默认方法", "不确定时使用此策略"),
        ("ABS_THRESHOLD", "绝对阈值判定", "ceiling_saturated (100%), floor_saturated (0)", "当KPI应该恒为100%或0%时"),
        ("RANGE_CONTEXT", "区间上下文判定 [μ±2σ]", "count_type (计数型)", "计数类KPI波动大时"),
        ("RELATIVE_CHANGE", "相对变化率判定", "ratio_type (比例型，非饱和)", "百分比类KPI非100%时"),
        ("ABS_DELTA_TREND", "绝对变化+Level Shift检测", "physical_stable (dB/dBm)", "物理量测量值"),
        ("QUANTILES", "分位数判定 (P50/P95)", "throughput_type (吞吐量)", "吞吐量数据"),
        ("TAIL_AWARE", "尾部感知判定 (P95/P99)", "latency_type (延迟)", "延迟/时间类数据"),
        ("WINDOW_COUNT", "滑窗计数判定", "low_frequency (低频事件)", "大部分为0的计数"),
        ("SEMANTIC_RULE", "语义规则判定", "config_type (配置型)", "配置开关状态"),
    ]
    
    for r_idx, (strategy, desc, suitable, when) in enumerate(ref_data, 4):
        ws_ref.cell(row=r_idx, column=1, value=strategy).border = thin_border
        ws_ref.cell(row=r_idx, column=2, value=desc).border = thin_border
        ws_ref.cell(row=r_idx, column=3, value=suitable).border = thin_border
        ws_ref.cell(row=r_idx, column=4, value=when).border = thin_border
    
    ws_ref.column_dimensions['A'].width = 20
    ws_ref.column_dimensions['B'].width = 35
    ws_ref.column_dimensions['C'].width = 45
    ws_ref.column_dimensions['D'].width = 35
    
    # Sheet 3: Quick Actions Summary
    ws_summary = wb.create_sheet(title="3_Quick_Summary")
    
    ws_summary['A1'] = "快速操作指南"
    ws_summary['A1'].font = Font(bold=True, size=14)
    
    summary_data = [
        "",
        f"📊 数据统计 (更新时间: {datetime.now().strftime('%Y-%m-%d %H:%M')})",
        f"  总KPI数量: {len(classification)}",
        f"  Other→Regression (需关注): {len(degraded_codes)} 个",
        f"  Regression→Other (已改善): {len(improved_codes)} 个",
        f"  用户已配置规则: {len(existing_rules)} 个",
        "",
        "🔧 推荐操作步骤:",
        "  1. 点击 Sheet '1_Custom_Rules'",
        "  2. 点击 'Change_Type' 列的筛选按钮",
        "  3. 选择 'Other→Regression ⚠️' 进行筛选",
        f"  4. 查看这 {len(degraded_codes)} 个KPI，决定是否需要修改策略",
        "  5. 如需修改：在 Override_Strategy 选择新策略，Enabled 设为 TRUE",
        "  6. 保存文件",
        "  7. 运行: python scripts/manage_custom_rules.py --apply",
        "",
        "💡 提示:",
        "  - 红色背景 = Other→Regression (可能需要改回 STANDARD_SPC)",
        "  - 绿色背景 = Regression→Other (新策略生效，通常不需要修改)",
        "  - 黄色背景 = 可编辑列",
        "  - 大部分情况下，只需关注红色行",
    ]
    
    for r_idx, text in enumerate(summary_data, 3):
        cell = ws_summary.cell(row=r_idx, column=1, value=text)
        if text.startswith("📊") or text.startswith("🔧") or text.startswith("💡"):
            cell.font = Font(bold=True)
    
    ws_summary.column_dimensions['A'].width = 80
    
    # Sheet 4: Modification History (placeholder)
    ws_history = wb.create_sheet(title="4_Modification_History")
    
    hist_headers = ['Timestamp', 'Code', 'KPI_Name', 'Original_Strategy', 'New_Strategy', 'Modified_By']
    for col, h in enumerate(hist_headers, 1):
        cell = ws_history.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    # Load existing history if available
    import json
    if HISTORY_FILE.exists():
        try:
            with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                history = json.load(f)
            for h_idx, record in enumerate(history[-100:], 2):
                ws_history.cell(row=h_idx, column=1, value=record.get('timestamp', '')).border = thin_border
                ws_history.cell(row=h_idx, column=2, value=record.get('code', '')).border = thin_border
                ws_history.cell(row=h_idx, column=3, value=record.get('kpi_name', '')).border = thin_border
                ws_history.cell(row=h_idx, column=4, value=record.get('original_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=5, value=record.get('new_strategy', '')).border = thin_border
                ws_history.cell(row=h_idx, column=6, value=record.get('modified_by', 'user')).border = thin_border
        except:
            pass
    
    hist_widths = [20, 18, 50, 18, 18, 15]
    for i, w in enumerate(hist_widths, 1):
        ws_history.column_dimensions[chr(64+i)].width = w
    
    ws_history.freeze_panes = 'A2'
    
    # Save
    wb.save(CUSTOM_RULES_FILE)
    
    print(f"\n✅ Custom rules Excel saved: {CUSTOM_RULES_FILE}")
    print(f"   Sheet 1: {len(classification)} KPIs (排序: 红色优先)")
    print(f"   Sheet 2: 策略说明")
    print(f"   Sheet 3: 快速操作指南")
    print(f"   Sheet 4: 修改历史")


def main():
    print("=" * 60)
    print("Sync Custom Rules from Classification")
    print("=" * 60)
    
    # Load data
    classification = load_classification_data()
    degraded_codes, improved_codes = load_comparison_data()
    existing_rules = load_existing_rules()
    
    # Create Excel
    create_custom_rules_excel(classification, degraded_codes, improved_codes, existing_rules)
    
    print("\n" + "=" * 60)
    print("下一步操作:")
    print("=" * 60)
    print(f"1. 打开: {CUSTOM_RULES_FILE}")
    print(f"2. 筛选 'Change_Type' = 'Other→Regression ⚠️' ({len(degraded_codes)}个)")
    print("3. 查看这些KPI，决定是否需要改为 STANDARD_SPC")
    print("4. 修改后运行: python scripts/manage_custom_rules.py --apply")


if __name__ == "__main__":
    main()
